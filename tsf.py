#!/usr/bin/env python3
"""TutorialStudyForge (TSF) - forge online tutorials into faithful study guides: PDF and editable Word (.docx).

Commands
    discover URL        list chapters found on a tutorial index page and write an editable chapters file
    build ...           fetch (or import) chapters, extract, typeset and verify a PDF (+ Markdown, manifest, report)
    inspect URL|FILE    show how one page is extracted (container, confidence, code blocks) to tune selectors
    selftest            build a small PDF from embedded sample pages to verify your installation

Design rules (see README.md): faithful compilation (no rewriting), byte-exact code, honest completeness
report, polite fetching (robots.txt, rate limit, cache), SSRF protection, never bypass blocks.

Requires: Python 3.10+, requests, beautifulsoup4, lxml, weasyprint, pypdf
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import html as htmllib
import ipaddress
import json
import logging
import os
import re
import socket
import sys
import time
import unicodedata
import urllib.parse
import urllib.robotparser
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

__version__ = "1.1.0"
SCHEMA_VERSION = 1
APP_NAME = "TutorialStudyForge"
APP_SHORT = "TSF"
USER_AGENT = f"TutorialStudyForge-TSF/{__version__} (personal-study compiler; honours robots.txt)"

log = logging.getLogger("t2p")


def say(msg: str = "") -> None:
    print(msg, file=sys.stderr, flush=True)


def _need(modname: str, pip: str | None = None):
    try:
        return __import__(modname)
    except ImportError:
        say(f"Missing dependency '{modname}'. Install with: pip install {pip or modname}")
        sys.exit(1)


requests = _need("requests")
_bs4 = _need("bs4", "beautifulsoup4 lxml")
BeautifulSoup, NavigableString, Tag, Comment = _bs4.BeautifulSoup, _bs4.NavigableString, _bs4.Tag, _bs4.Comment


# ======================================================================================
# Settings
# ======================================================================================
@dataclass
class Settings:
    out_dir: Path = Path("output")
    cache_dir: Path = Path.home() / ".cache" / "tsf"
    config_dir: Path = Path.home() / ".config" / "tsf"
    delay: float = 1.0                 # minimum seconds between requests to one host
    timeout: float = 30.0
    max_page_bytes: int = 10 * 1024 * 1024
    max_image_bytes: int = 3 * 1024 * 1024
    max_images_per_article: int = 40
    max_pages: int = 200
    retries: int = 2
    cache_ttl_hours: float = 24 * 7
    refresh: bool = False
    offline: bool = False
    allow_private: bool = False        # TESTING ONLY: disables SSRF protection
    ack_terms: list[str] = field(default_factory=list)
    min_words: int = 30
    images: bool = True
    code_cols: int = 90
    page_size: str = "A4"
    landscape: bool = False
    font_pt: float = 10.5
    margin_mm: float = 20.0
    cover: bool = True
    toc: bool = True
    toc_depth: int | None = None       # None = auto (chapters only when the guide has parts)
    chapter_breaks: bool = True
    docx: bool = False
    pdf: bool = True
    section_selector: str | None = None
    yes: bool = False
    source_urls: bool = True
    report_page: bool = True
    write_markdown: bool = False
    write_html: bool = False
    allow_partial: bool = False
    title: str = ""
    subtitle: str = ""
    author: str = ""
    selector: str | None = None
    drop_text: list[str] = field(default_factory=list)
    profiles_file: Path | None = None
    verbose: bool = False

    def public(self) -> dict:
        """Settings safe to record in the manifest (no local paths)."""
        keys = ["delay", "timeout", "min_words", "images", "code_cols", "page_size", "landscape", "font_pt",
                "margin_mm", "cover", "toc", "toc_depth", "source_urls", "report_page", "write_markdown",
                "write_html", "docx", "pdf", "max_pages", "selector", "drop_text"]
        return {k: getattr(self, k) for k in keys}


# ======================================================================================
# URL utilities and SSRF protection
# ======================================================================================
TRACKING_PARAMS = re.compile(r"^(utm_.*|fbclid|gclid|dclid|msclkid|mc_cid|mc_eid|igshid|ref|ref_src|_ga|yclid)$", re.I)


class UrlError(ValueError):
    pass


def normalize_url(url: str) -> str:
    """Normalise a URL for fetching/caching: lower-case scheme+host, drop fragment, default port, tracking params."""
    p = urllib.parse.urlsplit(url.strip())
    scheme = p.scheme.lower()
    host = (p.hostname or "").lower()
    port = p.port
    netloc = host
    if port and not ((scheme == "http" and port == 80) or (scheme == "https" and port == 443)):
        netloc = f"{host}:{port}"
    query = urllib.parse.urlencode(
        [(k, v) for k, v in urllib.parse.parse_qsl(p.query, keep_blank_values=True) if not TRACKING_PARAMS.match(k)]
    )
    path = p.path or "/"
    return urllib.parse.urlunsplit((scheme, netloc, path, query, ""))


def dedupe_key(url: str) -> str:
    """Key used to detect duplicate pages: normalised URL, www-less, no trailing slash, no index.html."""
    p = urllib.parse.urlsplit(normalize_url(url))
    host = p.netloc[4:] if p.netloc.startswith("www.") else p.netloc
    path = re.sub(r"/index\.html?$", "/", p.path)
    path = path.rstrip("/") or "/"
    return f"{host}{path}?{p.query}" if p.query else f"{host}{path}"


def validate_url_syntax(url: str, allow_private: bool = False) -> urllib.parse.SplitResult:
    try:
        p = urllib.parse.urlsplit(url.strip())
        port = p.port
    except ValueError as e:
        raise UrlError(f"Invalid URL: {e}") from e
    if p.scheme.lower() not in ("http", "https"):
        raise UrlError(f"Only http and https URLs are allowed (got '{p.scheme or 'none'}').")
    if not p.hostname:
        raise UrlError("URL has no host name.")
    if p.username or p.password:
        raise UrlError("URLs containing credentials are not allowed.")
    if port is not None and port not in (80, 443) and not allow_private:
        raise UrlError(f"Port {port} is not allowed (only 80 and 443).")
    if len(url) > 2048:
        raise UrlError("URL is too long.")
    return p


def _ip_is_public(ip: ipaddress._BaseAddress) -> bool:
    if isinstance(ip, ipaddress.IPv6Address) and ip.ipv4_mapped:
        ip = ip.ipv4_mapped
    if isinstance(ip, ipaddress.IPv6Address) and ip.sixtofour:
        ip = ip.sixtofour
    return ip.is_global and not ip.is_multicast


def assert_public_host(host: str, port: int) -> list[str]:
    """Resolve host and refuse private, loopback, link-local, reserved and metadata addresses."""
    try:
        infos = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except socket.gaierror as e:
        raise FetchError("dns", f"Could not resolve host '{host}': {e}") from e
    addrs = sorted({i[4][0] for i in infos})
    if not addrs:
        raise FetchError("dns", f"No addresses found for '{host}'.")
    for a in addrs:
        try:
            ip = ipaddress.ip_address(a.split("%")[0])
        except ValueError:
            raise FetchError("ssrf", f"Unparseable address for '{host}'.")
        if not _ip_is_public(ip):
            raise FetchError("ssrf", f"Refusing to connect to non-public address {a} for host '{host}'.")
    return addrs


def slugify(text: str, keep_case: bool = False, default: str = "tutorial") -> str:
    t = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    t = re.sub(r"[^A-Za-z0-9._ -]+", "", t).strip()
    t = re.sub(r"[\s_-]+", "-", t).strip("-.")
    if not keep_case:
        t = t.lower()
    return t[:80] or default


def safe_join(base: Path, *parts: str) -> Path:
    """Join paths and refuse anything that escapes `base`."""
    base_r = base.resolve()
    p = base_r.joinpath(*parts).resolve()
    if base_r != p and base_r not in p.parents:
        raise ValueError("Refusing to write outside the output directory.")
    return p


# ======================================================================================
# Fetching (the only place that touches the network)
# ======================================================================================
class FetchError(Exception):
    HINTS = {
        "access_denied": "The site refused automated access. Do not try to bypass it: save the page from your "
                         "browser and use 'build --html FILE_OR_DIR' instead.",
        "rate_limited": "The site is rate limiting you. Wait and retry later, or raise --delay.",
        "not_found": "The page does not exist (404/410). Remove it from the chapters file or fix the URL.",
        "robots": "robots.txt disallows automated access to this URL. Use 'build --html' with pages you saved yourself.",
        "ssrf": "The address is not a public internet address, so it was blocked.",
        "dns": "The host name could not be resolved. Check the URL and your connection.",
        "timeout": "The request timed out. Retry later or raise --timeout.",
        "network": "A network error occurred. Check your connection and retry.",
        "too_large": "The response was larger than the allowed limit.",
        "bad_type": "The response was not an HTML page.",
        "offline": "Offline mode is on and the page is not in the cache.",
        "host_blocked": "Fetching from this host was stopped after repeated 403/429 responses.",
        "http": "The server returned an error status.",
        "url": "The URL is not valid.",
    }

    def __init__(self, kind: str, message: str, status: int | None = None):
        super().__init__(message)
        self.kind, self.message, self.status = kind, message, status

    @property
    def hint(self) -> str:
        return self.HINTS.get(self.kind, "")


@dataclass
class Fetched:
    url: str            # the requested (normalised) URL
    final_url: str
    status: int
    body: bytes
    content_type: str
    fetched_at: str
    from_cache: bool


class Fetcher:
    def __init__(self, s: Settings):
        self.s = s
        self.sess = requests.Session()
        self.sess.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "en;q=0.9,*;q=0.5"})
        self.last_hit: dict[str, float] = {}
        self.robots: dict[str, urllib.robotparser.RobotFileParser | str] = {}
        self.denials: dict[str, int] = {}
        s.cache_dir.mkdir(parents=True, exist_ok=True)

    # ---- cache ----
    def _cache_paths(self, url: str) -> tuple[Path, Path]:
        h = hashlib.sha256(url.encode()).hexdigest()
        d = self.s.cache_dir / h[:2]
        return d / f"{h}.bin", d / f"{h}.json"

    def _cache_get(self, url: str) -> Fetched | None:
        if self.s.refresh:
            return None
        binp, jp = self._cache_paths(url)
        try:
            meta = json.loads(jp.read_text())
            body = binp.read_bytes()
            if hashlib.sha256(body).hexdigest() != meta["sha256"]:
                raise ValueError("cache checksum mismatch")
            age_h = (time.time() - meta["stored"]) / 3600
            if age_h > self.s.cache_ttl_hours and not self.s.offline:
                return None
            return Fetched(url, meta["final_url"], meta["status"], body, meta["content_type"], meta["fetched_at"], True)
        except FileNotFoundError:
            return None
        except Exception as e:  # corrupt cache entry: ignore and refetch
            log.warning("Ignoring corrupt cache entry for %s (%s)", url, e)
            return None

    def _cache_put(self, f: Fetched) -> None:
        binp, jp = self._cache_paths(f.url)
        try:
            binp.parent.mkdir(parents=True, exist_ok=True)
            binp.write_bytes(f.body)
            jp.write_text(json.dumps({
                "final_url": f.final_url, "status": f.status, "content_type": f.content_type,
                "fetched_at": f.fetched_at, "stored": time.time(), "sha256": hashlib.sha256(f.body).hexdigest(),
            }))
        except OSError as e:
            log.warning("Could not write cache: %s", e)

    # ---- politeness ----
    def _wait(self, host: str, extra_delay: float = 0.0) -> None:
        gap = max(self.s.delay, extra_delay)
        last = self.last_hit.get(host)
        if last is not None:
            remaining = gap - (time.monotonic() - last)
            if remaining > 0:
                time.sleep(remaining)
        self.last_hit[host] = time.monotonic()

    def _robots_for(self, p: urllib.parse.SplitResult):
        origin = f"{p.scheme}://{p.netloc}"
        if origin in self.robots:
            return self.robots[origin]
        rp: urllib.robotparser.RobotFileParser | str
        try:
            f = self._get_chain(f"{origin}/robots.txt", kind="robots")
            rp = urllib.robotparser.RobotFileParser()
            rp.parse(f.body.decode("utf-8", "replace").splitlines())
        except FetchError as e:
            if e.kind == "not_found":
                rp = "allow"                      # no robots.txt: everything allowed
            elif e.kind in ("access_denied",):
                rp = "deny"                       # 401/403 on robots.txt: be conservative (stdlib semantics)
            elif e.kind in ("ssrf", "dns", "offline"):
                raise
            else:
                rp = "deny"                       # server error / unreachable: be conservative
                log.warning("robots.txt unavailable for %s (%s); treating the site as disallowed.", origin, e.message)
        self.robots[origin] = rp
        return rp

    def robots_allows(self, url: str) -> bool:
        p = urllib.parse.urlsplit(url)
        rp = self._robots_for(p)
        if rp == "allow":
            return True
        if rp == "deny":
            return False
        return rp.can_fetch(USER_AGENT, url)

    # ---- main entry ----
    def get(self, url: str, kind: str = "page") -> Fetched:
        """Fetch a page or image with every protection applied. kind: 'page' | 'image'."""
        try:
            validate_url_syntax(url, self.s.allow_private)
        except UrlError as e:
            raise FetchError("url", str(e)) from e
        return self._get_chain(normalize_url(url), kind=kind)

    def _get_chain(self, url: str, kind: str) -> Fetched:
        cached = self._cache_get(url)
        if cached:
            return cached
        if self.s.offline:
            raise FetchError("offline", f"{url} is not in the cache.")
        orig = url
        for _hop in range(6):
            p = validate_url_syntax(url, self.s.allow_private)
            host = p.hostname or ""
            port = p.port or (443 if p.scheme == "https" else 80)
            if self.denials.get(host, 0) >= 3:
                raise FetchError("host_blocked", f"Stopped requesting {host} after repeated 403/429 responses.")
            if not self.s.allow_private:
                assert_public_host(host, port)
            crawl_delay = 0.0
            if kind != "robots":
                if not self.robots_allows(url):
                    raise FetchError("robots", f"robots.txt disallows fetching {url}")
                rp = self.robots.get(f"{p.scheme}://{p.netloc}")
                if isinstance(rp, urllib.robotparser.RobotFileParser):
                    crawl_delay = float(rp.crawl_delay(USER_AGENT) or 0)
            resp = self._request(url, host, kind, crawl_delay)
            if resp.status_code in (301, 302, 303, 307, 308):
                loc = resp.headers.get("Location")
                resp.close()
                if not loc:
                    raise FetchError("http", "Redirect without Location header.", resp.status_code)
                url = normalize_url(urllib.parse.urljoin(url, loc))
                continue
            try:
                body = self._read_body(resp, kind)
            finally:
                resp.close()
            f = Fetched(orig, url, resp.status_code, body, resp.headers.get("Content-Type", ""),
                        datetime.now(timezone.utc).isoformat(timespec="seconds"), False)
            self._cache_put(f)
            return f
        raise FetchError("http", "Too many redirects.")

    def _request(self, url: str, host: str, kind: str, crawl_delay: float):
        last_err: FetchError | None = None
        for attempt in range(self.s.retries + 1):
            self._wait(host, crawl_delay)
            try:
                resp = self.sess.get(url, stream=True, allow_redirects=False,
                                     timeout=(10, self.s.timeout))
            except requests.Timeout:
                last_err = FetchError("timeout", f"Timed out fetching {url}")
            except requests.ConnectionError as e:
                last_err = FetchError("network", f"Connection error for {url}: {e.__class__.__name__}")
            except requests.RequestException as e:
                raise FetchError("network", f"Request failed: {e.__class__.__name__}") from e
            else:
                sc = resp.status_code
                if sc in (404, 410):
                    resp.close()
                    raise FetchError("not_found", f"HTTP {sc} for {url}", sc)
                if sc in (401, 403):
                    resp.close()
                    self.denials[host] = self.denials.get(host, 0) + 1
                    raise FetchError("access_denied", f"HTTP {sc} for {url}", sc)
                if sc == 429:
                    self.denials[host] = self.denials.get(host, 0) + 1
                    wait = _retry_after(resp.headers.get("Retry-After"))
                    resp.close()
                    if attempt < self.s.retries and wait is not None and wait <= 60:
                        say(f"  rate limited by {host}; waiting {wait:.0f}s")
                        time.sleep(wait)
                        last_err = FetchError("rate_limited", f"HTTP 429 for {url}", 429)
                        continue
                    raise FetchError("rate_limited", f"HTTP 429 for {url}", 429)
                if 500 <= sc < 600:
                    resp.close()
                    last_err = FetchError("http", f"HTTP {sc} for {url}", sc)
                elif sc >= 400:
                    resp.close()
                    raise FetchError("http", f"HTTP {sc} for {url}", sc)
                else:
                    return resp
            if attempt < self.s.retries:
                time.sleep(min(2 ** attempt, 8))          # bounded backoff, transient errors only
        assert last_err is not None
        raise last_err

    def _read_body(self, resp, kind: str) -> bytes:
        ctype = (resp.headers.get("Content-Type") or "").lower()
        limit = self.s.max_image_bytes if kind == "image" else self.s.max_page_bytes
        if kind == "page" and ctype and not re.search(r"html|xml|text/plain", ctype):
            raise FetchError("bad_type", f"Unexpected content type '{ctype}'.")
        cl = resp.headers.get("Content-Length")
        if cl and cl.isdigit() and int(cl) > limit:
            raise FetchError("too_large", f"Response is {int(cl)} bytes (limit {limit}).")
        chunks, total = [], 0
        try:
            for chunk in resp.iter_content(65536):
                total += len(chunk)
                if total > limit:
                    raise FetchError("too_large", f"Response exceeded {limit} bytes.")
                chunks.append(chunk)
        except requests.RequestException as e:
            raise FetchError("network", f"Download interrupted: {e.__class__.__name__}") from e
        return b"".join(chunks)


def _retry_after(v: str | None) -> float | None:
    if not v:
        return 5.0
    if v.isdigit():
        return float(v)
    return None


# ======================================================================================
# Terms acknowledgement (one-time, per host)
# ======================================================================================
def host_acknowledged(s: Settings, host: str) -> bool:
    norm = lambda h: (h[4:] if h.startswith('www.') else h).lower()
    if s.allow_private or norm(host) in {norm(fix_host(x)) for x in s.ack_terms}:
        return True
    f = s.config_dir / "acknowledged_hosts.json"
    try:
        known = json.loads(f.read_text())
    except Exception:
        known = []
    if norm(host) in {norm(k) for k in known}:
        return True
    if sys.stdin.isatty():
        say(f"\nBefore fetching from {host}: this tool copies third-party content for PERSONAL study only.\n"
            f"Check that {host}'s terms of use allow automated access, and do not share the output.")
        if input(f"Have you reviewed {host}'s terms and want to continue? [y/N] ").strip().lower() == "y":
            try:
                s.config_dir.mkdir(parents=True, exist_ok=True)
                f.write_text(json.dumps(sorted(set(known + [host]))))
            except OSError:
                pass
            return True
        return False
    say(f"Refusing to fetch from {host} non-interactively without --ack-terms {host} "
        f"(confirms you reviewed its terms of use).")
    return False


# ======================================================================================
# Site profiles (selectors are configuration, not code)
# ======================================================================================
_SPHINX = {
    "name": "Sphinx / Read the Docs documentation",
    "content": ["div[role=main]", "article.bd-article", "div.body", "main", "div.document"],
    "toc": ["div.toctree-wrapper", "nav.bd-links", "div.sphinxsidebarwrapper", "div.sidebar-tree"],
    "remove": ["div.sphinxsidebar", "div.related", ".rst-footer-buttons", ".wy-nav-top", "div.admonition-todo"],
    "code_selectors": [],
}

# Built-in rules are written from the sites' usual page layouts. Sites change their HTML without notice, so treat
# them as a head start: `discover` shows what was found, `inspect` shows how one page is read, and the --selector /
# --section-selector / --link-selector / --profiles options fix anything that is off. The generic detectors
# (article/main detection, sidebar detection, heading-based sections) work for most other sites.
DEFAULT_PROFILES: dict[str, dict] = {
    "geeksforgeeks.org": {
        "name": "GeeksforGeeks", "example": "https://www.geeksforgeeks.org/sql-tutorial/",
        "content": ["div.article--viewer_content", "article", "div.entry-content", "main"],
        "toc": [], "remove": [], "code_selectors": [],
    },
    "w3schools.com": {
        "name": "W3Schools", "example": "https://www.w3schools.com/sql/default.asp",
        "content": ["#main"], "toc": ["#leftmenuinner"],
        "remove": ["a.w3-btn", "a.ws-btn", "a.ws-green", "a.ws-black", ".nextprev", "#mainLeaderboard",
                   "#midcontentadcontainer", "#bottom_ad", ".ws-ad", "[id^=snigel]", "[id^=adngin]",
                   "#google_translate_element"],
        "code_selectors": ["div.w3-code"],
    },
    "tutorialspoint.com": {
        "name": "TutorialsPoint", "example": "https://www.tutorialspoint.com/sql/index.htm",
        "content": ["div.tutorial-content", "#mainContent", "main"],
        "toc": ["ul.toc.chapters", "ul.chapters", "#sidebar", "aside"],
        "remove": [".pre-btn", ".nxt-btn", ".btn-group", "div.prev-next", "#google_translate_element"],
        "code_selectors": [],
    },
    "javatpoint.com": {
        "name": "JavaTpoint", "example": "https://www.javatpoint.com/java-tutorial",
        "content": ["#city", "div#content"], "toc": ["div.leftmenu", "#menu", "nav.leftmenu"],
        "remove": ["div.nexttopicdiv", ".nexttopic", "#bottomnext", "div.preview", "#bottom"],
        "code_selectors": [],
    },
    "programiz.com": {
        "name": "Programiz", "example": "https://www.programiz.com/python-programming",
        "content": ["div.editor-contents", "article", "main"], "toc": [], "remove": [], "code_selectors": [],
    },
    "developer.mozilla.org": {
        "name": "MDN Web Docs", "example": "https://developer.mozilla.org/en-US/docs/Learn_web_development",
        "content": ["article.main-page-content", "main article", "article"],
        "toc": ["nav.sidebar", "#sidebar-quicklinks", ".sidebar-body"],
        "remove": [".copy-icon", ".code-example .example-header", ".play-button", "section.baseline-indicator"],
        "code_selectors": [],
    },
    "docs.python.org": {**_SPHINX, "name": "Python documentation", "example": "https://docs.python.org/3/tutorial/index.html"},
    "readthedocs.io": {**_SPHINX, "example": "https://<project>.readthedocs.io/en/latest/"},
    "sphinx-doc.org": dict(_SPHINX),
    "realpython.com": {
        "name": "Real Python", "example": "https://realpython.com/python-basic-syntax/",
        "content": ["div.article-body", "article", "main"], "toc": [], "remove": [".sidebar-module"],
        "code_selectors": [],
    },
    "baeldung.com": {
        "name": "Baeldung", "example": "https://www.baeldung.com/java-streams",
        "content": ["div.post-content", "article", ".entry-content"], "toc": [], "remove": [], "code_selectors": [],
    },
    "docs.oracle.com": {
        "name": "Oracle Java Tutorials", "example": "https://docs.oracle.com/javase/tutorial/",
        "content": ["#PageContent", "div.PageContent", "#MainFlow", "main"],
        "toc": ["#TableOfContents", "div.TOC"], "remove": [".NavBit", "#BreadCrumbs", ".breadcrumbs"],
        "code_selectors": [],
    },
    "*": {"name": "any other site",
          "content": ["article", "div[role=main]", "main", "[role=main]", ".markdown-body", "div.md-content",
                      "div#content", "div.content"],
          "toc": [], "remove": [], "code_selectors": []},
}

GENERIC_REMOVE = ["a.headerlink", "a.anchor", "a.hash-link", ".copy-button", ".copybtn", ".ad", ".ads",
                  ".adsbygoogle", "#disqus_thread"]


def load_profile(host: str, s: Settings) -> dict:
    profiles = copy.deepcopy(DEFAULT_PROFILES)
    if s.profiles_file:
        try:
            user = json.loads(Path(s.profiles_file).read_text())
            for k, v in user.items():
                profiles.setdefault(k, {}).update(v)
        except Exception as e:
            say(f"Could not read profiles file: {e}")
            sys.exit(2)
    h = host.lower()
    h = h[4:] if h.startswith("www.") else h
    prof = profiles["*"]
    for key, cand in profiles.items():
        if key != "*" and (h == key or h.endswith("." + key)):
            prof = {**profiles["*"], **cand}
            break
    prof = dict(prof)
    prof["remove"] = GENERIC_REMOVE + list(prof.get("remove", []))
    return prof


# ======================================================================================
# HTML extraction -> typed block model
# ======================================================================================
DROP_TAGS = {"script", "style", "noscript", "iframe", "form", "button", "input", "select", "textarea", "svg",
             "canvas", "audio", "video", "object", "embed", "template", "dialog", "link", "meta", "head", "title",
             "nav", "map", "datalist", "label"}
INLINE_TAGS = {"a", "abbr", "b", "bdi", "bdo", "big", "cite", "code", "data", "dfn", "em", "i", "kbd", "mark", "q",
               "s", "samp", "small", "span", "strong", "sub", "sup", "time", "tt", "u", "var", "wbr", "br", "img",
               "font", "del", "ins", "math"}
BLOCK_TAGS = ["p", "pre", "table", "ul", "ol", "div", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "figure",
              "hr", "section", "article", "dl", "details", "li", "main", "aside", "header", "footer"]
JUNK_TOKEN = re.compile(
    r"^(ad|ads|advert\w*|adsbygoogle|adunit\w*|banner\w*|cookie\w*|consent|share\w*|social\w*|newsletter|subscribe\w*|"
    r"promo\w*|popup|modal|sponsor\w*|breadcrumb\w*|related\w*|recommend\w*|comment\w*|disqus\w*|sticky|floating|"
    r"toolbar|pagination|copy\w*|btn|button|rating|vote|trending|widget)$", re.I)
CALLOUT_TOKENS = {"note", "tip", "warning", "info", "caution", "important", "callout", "admonition", "alert", "hint",
                  "danger", "notice"}
ACCESS_PAGE = re.compile(r"(access denied|just a moment|enable javascript and cookies|verify you are human|captcha|"
                         r"are you a robot|unusual traffic|attention required|403 forbidden|request blocked)", re.I)
DEFAULT_DROP_TEXT = [r"^(advertisement|sponsored|skip to content|trending now|improve|suggest changes|like article|"
                     r"save article|share your thoughts in the comments)$"]
GUTTER_SEL = (".linenos, .line-numbers, .line-number, .linenumber, .lineno, .gutter, .line-numbers-rows, "
              "[data-line-number], td.gutter, .hljs-ln-numbers")
SKIP_IMG_TOKENS = re.compile(r"(emoji|icon|avatar|logo|sprite|pixel|tracking|badge|spinner)", re.I)


def classes(el: Tag) -> list[str]:
    c = el.get("class") or []
    return c if isinstance(c, list) else str(c).split()


def tokens(el: Tag) -> list[str]:
    out = []
    for v in classes(el) + ([el.get("id")] if el.get("id") else []):
        out += [t for t in re.split(r"[^A-Za-z0-9]+", str(v)) if t]
    return out


def collapse_ws(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def esc(s: str) -> str:
    return htmllib.escape(s, quote=False)


def abs_url(href: str | None, base: str | None) -> str | None:
    if not href:
        return None
    href = href.strip()
    if href.startswith(("#", "javascript:", "mailto:", "tel:", "data:")):
        return None
    u = urllib.parse.urljoin(base, href) if base else href
    return u if u.lower().startswith(("http://", "https://")) else None


class Extractor:
    """Turns a cleaned HTML container into typed blocks. Pure function of the HTML: no network."""

    def __init__(self, base_url: str | None, profile: dict, s: Settings):
        self.base = base_url
        self.profile = profile
        self.s = s
        self.warnings: list[str] = []
        self.code_ids: set[int] = set()

    # ---- inline ----
    def inline(self, nodes) -> str:
        out: list[str] = []
        for ch in nodes:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                out.append(esc(str(ch)))
                continue
            if not isinstance(ch, Tag):
                continue
            n = ch.name
            if n in DROP_TAGS:
                continue
            if n == "br":
                out.append("<br>")
            elif n == "img":
                continue
            elif n == "a":
                inner = self.inline(ch.children)
                href = abs_url(ch.get("href"), self.base)
                out.append(f'<a href="{htmllib.escape(href, quote=True)}">{inner}</a>' if href and inner.strip() else inner)
            elif n in ("strong", "b"):
                inner = self.inline(ch.children)
                out.append(f"<strong>{inner}</strong>" if inner.strip() else inner)
            elif n in ("em", "i", "cite", "dfn"):
                inner = self.inline(ch.children)
                out.append(f"<em>{inner}</em>" if inner.strip() else inner)
            elif n in ("code", "kbd", "samp", "tt", "var"):
                txt = code_text(ch).replace("\n", " ")
                out.append(f"<code>{esc(txt)}</code>" if txt.strip() else "")
            elif n in ("sub", "sup"):
                out.append(f"<{n}>{self.inline(ch.children)}</{n}>")
            elif n == "math":
                tex = ch.find("annotation", attrs={"encoding": re.compile("tex", re.I)})
                txt = tex.get_text() if tex else ch.get_text()
                out.append(f"<code>{esc(collapse_ws(txt))}</code>")
            else:
                out.append(self.inline(ch.children))
        return "".join(out)

    @staticmethod
    def tidy(html_str: str) -> str:
        s = html_str.replace("\xa0", " ")
        s = re.sub(r"[ \t\r\n]+", " ", s)
        s = re.sub(r"\s*<br>\s*", "<br>", s)
        s = re.sub(r"^(<br>)+|(<br>)+$", "", s.strip())
        return s.strip()

    # ---- images ----
    def image_block(self, img: Tag, caption: str = "") -> dict | None:
        src = img.get("data-src") or img.get("data-lazy-src") or img.get("data-original") or img.get("src") or ""
        if not src and img.get("srcset"):
            src = img["srcset"].split(",")[0].strip().split(" ")[0]
        if not src or src.startswith("data:"):
            return None
        if SKIP_IMG_TOKENS.search(" ".join(tokens(img)) + " " + src):
            return None
        for dim in ("width", "height"):
            v = str(img.get(dim, "")).strip()
            if v.isdigit() and int(v) <= 24:
                return None
        return {"type": "image", "src": src, "url": abs_url(src, self.base) if self.base else None,
                "alt": collapse_ws(img.get("alt", "")), "caption": caption, "asset": None}

    # ---- block walk ----
    def walk(self, el: Tag) -> list[dict]:
        out: list[dict] = []
        buf: list = []

        def flush():
            nonlocal buf
            nodes, buf = buf, []
            if not nodes:
                return
            tmp = BeautifulSoup("<x></x>", "html.parser").x
            for n in nodes:
                tmp.append(copy.copy(n))
            imgs = tmp.find_all("img")
            html_s = self.tidy(self.inline(tmp.children))
            if re.sub(r"<[^>]+>|\s", "", html_s):
                out.append({"type": "paragraph", "html": html_s})
            for im in imgs:
                b = self.image_block(im)
                if b:
                    out.append(b)

        for ch in el.children:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                if str(ch).strip() or buf:
                    buf.append(ch)
                continue
            if not isinstance(ch, Tag) or ch.name in DROP_TAGS:
                continue
            if ch.name in INLINE_TAGS and id(ch) not in self.code_ids and not ch.find(BLOCK_TAGS):
                buf.append(ch)
                continue
            flush()
            self.block(ch, out)
        flush()
        return out

    def block(self, el: Tag, out: list[dict]) -> None:
        n = el.name
        if id(el) in self.code_ids:
            self._code(el, out)
        elif re.fullmatch(r"h[1-6]", n):
            h = self.tidy(self.inline(el.children))
            if re.sub(r"<[^>]+>|\s", "", h):
                out.append({"type": "heading", "level": int(n[1]), "html": h})
        elif n == "pre":
            self._code(el, out)
        elif n in ("ul", "ol"):
            out.append(self._list(el))
        elif n == "table":
            b = self._table(el)
            if b:
                out.append(b)
        elif n == "blockquote":
            inner = self.walk(el)
            if inner:
                out.append({"type": "quote", "kind": "", "blocks": inner})
        elif n == "figure":
            cap_el = el.find("figcaption")
            cap = collapse_ws(cap_el.get_text()) if cap_el else ""
            if cap_el:
                cap_el.decompose()
            inner = self.walk(el)
            for b in inner:
                if b["type"] == "image" and cap and not b["caption"]:
                    b["caption"] = cap
                    cap = ""
            out.extend(inner)
        elif n == "hr":
            out.append({"type": "hr"})
        elif n == "dl":
            for c in el.find_all(["dt", "dd"], recursive=False):
                h = self.tidy(self.inline(c.children))
                if h:
                    out.append({"type": "paragraph", "html": f"<strong>{h}</strong>" if c.name == "dt" else h})
        elif n == "math" or (set(tokens(el)) & {"katex", "mathjax", "MathJax"}):
            ann = el.find("annotation", attrs={"encoding": re.compile("tex", re.I)})
            tex = (ann.get_text() if ann else el.get_text()).strip()
            if tex:
                if not ann:
                    self.warnings.append("Equation had no TeX source; used its text content.")
                out.append({"type": "equation", "tex": tex})
        elif set(t.lower() for t in classes(el)) & CALLOUT_TOKENS and not el.find("pre", recursive=False):
            inner = self.walk(el)
            kind = sorted(set(t.lower() for t in classes(el)) & CALLOUT_TOKENS)[0]
            if inner:
                out.append({"type": "quote", "kind": kind, "blocks": inner})
        else:
            out.extend(self.walk(el))

    def _code(self, el: Tag, out: list[dict]) -> None:
        text = code_text(el)
        if not text.strip():
            return
        lang = code_language(el)
        out.append({"type": "code", "language": lang, "text": text})

    def _list(self, el: Tag) -> dict:
        items = []
        for li in el.find_all("li", recursive=False):
            blocks = self.walk(li)
            if blocks:
                items.append(blocks)
        start = el.get("start")
        return {"type": "list", "ordered": el.name == "ol", "start": int(start) if str(start).isdigit() else 1,
                "items": items}

    def _table(self, el: Tag) -> dict | None:
        if el.find("pre") or el.select_one(GUTTER_SEL):
            for g in el.select(GUTTER_SEL):
                g.decompose()
            pre = el.find("pre")
            text = code_text(pre if pre else el)
            if text.strip():
                return {"type": "code", "language": code_language(pre or el), "text": text}
            return None
        rows = []
        for tr in el.find_all("tr"):
            if tr.find_parent("table") is not el:
                continue
            cells = []
            for td in tr.find_all(["td", "th"], recursive=False):
                if td.find("table"):
                    self.warnings.append("A nested table was flattened to text.")
                h = self.tidy(self.inline(td.children))
                span = lambda k: int(td.get(k)) if str(td.get(k, "")).isdigit() and 1 < int(td.get(k)) < 50 else 1
                cells.append({"html": h, "th": td.name == "th", "colspan": span("colspan"), "rowspan": span("rowspan")})
            if cells:
                head = tr.find_parent("thead") is not None or all(c["th"] for c in cells)
                rows.append({"head": head, "cells": cells})
        if not rows:
            return None
        cap = el.find("caption")
        return {"type": "table", "caption": collapse_ws(cap.get_text()) if cap else "", "rows": rows}


def code_language(el: Tag) -> str:
    cand: list[str] = []
    chain = [el] + list(el.find_all(["code"], limit=1)) + list(el.parents)[:3]
    for e in chain:
        if not isinstance(e, Tag):
            continue
        cand += classes(e)
        for k in ("data-lang", "data-language", "lang"):
            if e.get(k):
                cand.append(f"language-{e.get(k)}")
    for c in cand:
        m = re.match(r"^(?:language|lang|highlight-source|highlight|brush:?|prettyprint\s*lang)[-_:]?(.+)$", c, re.I)
        if m and re.fullmatch(r"[A-Za-z0-9#+.\-_]{1,20}", m.group(1)) and m.group(1).lower() not in (
                "wrapper", "code", "pre", "none", "text-plain"):
            return m.group(1).lower()
    return ""


def code_text(el: Tag) -> str:
    """Extract code exactly: keep whitespace, turn <br>/line elements into newlines, drop gutters and buttons."""
    el = copy.copy(el)
    for g in el.select(GUTTER_SEL):
        g.decompose()
    for b in el.find_all(["button", "script", "style", "svg"]):
        b.decompose()
    parts: list[str] = []

    def ends_nl() -> bool:
        return bool(parts) and parts[-1].endswith("\n")

    def rec(node: Tag) -> None:
        for ch in node.children:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                parts.append(str(ch))
            elif isinstance(ch, Tag):
                if ch.name == "br":
                    parts.append("\n")
                    continue
                rec(ch)
                is_line = ch.name in ("div", "li", "p", "tr") or any(c in ("line", "code-line", "cl") for c in classes(ch))
                if is_line and not ends_nl():
                    nxt = ch.next_sibling
                    if not (isinstance(nxt, NavigableString) and str(nxt).startswith("\n")):
                        parts.append("\n")

    rec(el)
    text = "".join(parts).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ").replace("\u200b", "")
    return text.strip("\n").rstrip()


# ---- container finding / cleaning ----
def link_density(el: Tag) -> float:
    total = len(el.get_text(" ", strip=True)) or 1
    links = sum(len(a.get_text(" ", strip=True)) for a in el.find_all("a"))
    return links / total


def word_count(el: Tag) -> int:
    return len(el.get_text(" ", strip=True).split())


def find_container(soup: BeautifulSoup, profile: dict, override: str | None) -> tuple[Tag | None, str, float]:
    """Return (container, method, penalty). Method is 'selector:<css>' or 'heuristic'."""
    sels = [override] if override else profile.get("content", [])
    for sel in sels:
        try:
            el = soup.select_one(sel)
        except Exception:
            continue
        if el is not None and word_count(el) >= 30:
            return el, f"selector:{sel}", 0.0 if not override else 0.0
    # Heuristic: score paragraph/pre text, propagate to ancestors, descend into the dominant child.
    scores: dict[int, float] = {}
    nodes: dict[int, Tag] = {}
    for blk in soup.find_all(["p", "pre", "li", "td"]):
        if blk.find_parent(["nav", "footer", "header", "aside"]):
            continue
        txt = blk.get_text(" ", strip=True)
        if len(txt) < 25 and blk.name != "pre":
            continue
        sc = len(txt) * (2 if blk.name == "pre" else 1) - 2 * sum(len(a.get_text()) for a in blk.find_all("a"))
        node = blk.parent
        depth = 0
        while isinstance(node, Tag) and node.name not in ("[document]", "html") and depth < 25:
            scores[id(node)] = scores.get(id(node), 0) + sc * (0.97 ** depth)
            nodes[id(node)] = node
            node = node.parent
            depth += 1
    if not scores:
        return (soup.body if soup.body else None), "heuristic:body", 0.25
    best = max(scores, key=scores.get)
    cur = nodes[best]
    while True:
        kids = [c for c in cur.children if isinstance(c, Tag) and id(c) in scores]
        top = max(kids, key=lambda k: scores[id(k)], default=None)
        if top is not None and scores[id(top)] >= 0.85 * scores[id(cur)]:
            cur = top
        else:
            break
    return cur, "heuristic", 0.15


def clean_container(cont: Tag, profile: dict, s: Settings, warnings: list[str]) -> Tag:
    cont = copy.copy(cont)
    for sel in profile.get("remove", []):
        try:
            for e in cont.select(sel):
                e.decompose()
        except Exception:
            pass
    for e in cont.find_all(DROP_TAGS - {"label", "map", "head", "title", "meta", "link"}):
        e.decompose()
    for e in cont.find_all(["nav", "header", "footer"]):
        e.decompose()
    for e in list(cont.find_all(True)):
        if e.decomposed if hasattr(e, "decomposed") else False:
            continue
        if e.get("hidden") is not None or e.get("aria-hidden") == "true" and not e.get_text(strip=True):
            e.decompose()
            continue
        style = (e.get("style") or "").replace(" ", "").lower()
        if "display:none" in style or "visibility:hidden" in style:
            e.decompose()
    removed_kept: list[str] = []
    for e in list(cont.find_all(True)):
        if getattr(e, "decomposed", False) or e is cont or e.parent is None:
            continue
        if any(JUNK_TOKEN.match(t) for t in tokens(e)) or e.name == "aside":
            # Only drop it if it is not meaningful learning material (no code/tables, short or link-heavy).
            if e.find(["pre", "table"]) or e.name == "pre":
                removed_kept.append(" ".join(classes(e))[:40])
                continue
            if link_density(e) > 0.35 or word_count(e) < 60 or not e.find("p"):
                e.decompose()
            else:
                removed_kept.append(" ".join(classes(e))[:40])
    if removed_kept:
        warnings.append(f"Kept {len(removed_kept)} element(s) matching ad/widget patterns because they look like content "
                        f"(classes: {', '.join(sorted(set(removed_kept))[:4])}).")
    return cont


def plain_text(blocks: list[dict]) -> str:
    out: list[str] = []

    def rec(bs):
        for b in bs:
            t = b["type"]
            if t in ("paragraph", "heading"):
                out.append(re.sub(r"<[^>]+>", "", htmllib.unescape(b["html"])))
            elif t == "code":
                out.append(b["text"])
            elif t == "equation":
                out.append(b["tex"])
            elif t == "list":
                for it in b["items"]:
                    rec(it)
            elif t == "quote":
                rec(b["blocks"])
            elif t == "table":
                for r in b["rows"]:
                    out.extend(re.sub(r"<[^>]+>", "", htmllib.unescape(c["html"])) for c in r["cells"])
            elif t == "image":
                out.append(b["alt"] + " " + b["caption"])
    rec(blocks)
    return "\n".join(out)


def count_blocks(blocks: list[dict]) -> dict:
    c = {"headings": 0, "paragraphs": 0, "code_blocks": 0, "tables": 0, "lists": 0, "images": 0, "equations": 0}
    key = {"heading": "headings", "paragraph": "paragraphs", "code": "code_blocks", "table": "tables", "list": "lists",
           "image": "images", "equation": "equations"}

    def rec(bs):
        for b in bs:
            if b["type"] in key:
                c[key[b["type"]]] += 1
            if b["type"] == "list":
                for it in b["items"]:
                    rec(it)
            elif b["type"] == "quote":
                rec(b["blocks"])
    rec(blocks)
    return c


def normalise_headings(blocks: list[dict], title: str) -> list[dict]:
    """Drop a leading heading equal to the chapter title; shift levels so the shallowest becomes h2."""
    t_norm = collapse_ws(title).lower()
    out = list(blocks)
    while out and out[0]["type"] == "heading" and collapse_ws(re.sub(r"<[^>]+>", "", out[0]["html"])).lower() == t_norm:
        out.pop(0)
    levels = [b["level"] for b in out if b["type"] == "heading"]
    if levels:
        shift = 2 - min(levels)
        for b in out:
            if b["type"] == "heading":
                b["level"] = max(2, min(6, b["level"] + shift))
    dedup: list[dict] = []
    for b in out:
        if dedup and b["type"] == "heading" and dedup[-1]["type"] == "heading" and dedup[-1]["html"] == b["html"]:
            continue
        dedup.append(b)
    return dedup


def drop_furniture_text(blocks: list[dict], patterns: list[str]) -> list[dict]:
    rx = [re.compile(p, re.I) for p in patterns]
    keep = []
    for b in blocks:
        if b["type"] == "paragraph":
            txt = collapse_ws(re.sub(r"<[^>]+>", "", htmllib.unescape(b["html"])))
            if any(r.search(txt) for r in rx):
                continue
        keep.append(b)
    return keep


def page_title(soup: BeautifulSoup, cont: Tag | None) -> str:
    for scope in (cont, soup):
        if scope is not None:
            h1 = scope.find("h1")
            if h1 and collapse_ws(h1.get_text()):
                return collapse_ws(h1.get_text())
    og = soup.find("meta", attrs={"property": "og:title"})
    if og and og.get("content"):
        return collapse_ws(og["content"])
    if soup.title and soup.title.string:
        return collapse_ws(soup.title.string)
    return ""


def extract_article(data: bytes, url: str, base_url: str | None, s: Settings, profile_host: str | None = None) -> dict:
    """Extract one article. Never raises for content problems: returns status/error instead."""
    host = profile_host or (urllib.parse.urlsplit(url).hostname or "")
    profile = load_profile(host, s)
    soup = BeautifulSoup(data, "lxml")
    art: dict = {"title": "", "blocks": [], "warnings": [], "method": "", "confidence": 0.0, "status": "failed",
                 "error": None, "canonical_url": None, "words": 0, "content_hash": "", "counts": {}}
    canon = soup.find("link", rel="canonical")
    if canon and canon.get("href"):
        cu = abs_url(canon["href"], base_url or url)
        if cu:
            art["canonical_url"] = normalize_url(cu)
    cont, method, penalty = find_container(soup, profile, s.selector)
    art["method"] = method
    if cont is None:
        art["error"] = {"kind": "extract", "message": "No article content found.",
                        "hint": "Run 'inspect' on the page; the layout may have changed or content may need JavaScript."}
        return art
    art["title"] = page_title(soup, cont)
    warnings: list[str] = art["warnings"]
    cleaned = clean_container(cont, profile, s, warnings)
    ex = Extractor(base_url, profile, s)
    for sel in profile.get("code_selectors", []):
        try:
            ex.code_ids |= {id(e) for e in cleaned.select(sel)}
        except Exception:
            pass
    blocks = ex.walk(cleaned)
    warnings += ex.warnings
    blocks = drop_furniture_text(blocks, DEFAULT_DROP_TEXT + s.drop_text)
    blocks = normalise_headings(blocks, art["title"])
    art["blocks"] = blocks
    text = plain_text(blocks)
    words = len(text.split())
    art["words"] = words
    art["counts"] = count_blocks(blocks)
    art["content_hash"] = hashlib.sha256(re.sub(r"\W+", " ", text.lower()).strip().encode()).hexdigest()[:20]
    # ---- validation & confidence ----
    conf = 1.0 - penalty
    page_txt = soup.get_text(" ", strip=True)
    if words < 150 and ACCESS_PAGE.search(page_txt[:3000]):
        art["error"] = {"kind": "access_page", "message": "The page looks like an access-denied, captcha or "
                        "bot-check page, not an article.",
                        "hint": FetchError.HINTS["access_denied"]}
        art["confidence"] = 0.0
        return art
    if not art["title"]:
        conf -= 0.3
        warnings.append("No title found.")
    if words < s.min_words:
        art["error"] = {"kind": "empty", "message": f"Only {words} words extracted (minimum {s.min_words}).",
                        "hint": "The article may require JavaScript or the layout changed. Run 'inspect' and adjust "
                                "--selector or the profile."}
        art["confidence"] = 0.0
        return art
    if words < 150:
        conf -= 0.2
        warnings.append(f"Short article ({words} words).")
    if art["counts"]["paragraphs"] == 0:
        conf -= 0.3
        warnings.append("No paragraphs detected.")
    ld = link_density(cleaned)
    if ld > 0.5:
        conf -= 0.3
        warnings.append(f"High link density ({ld:.0%}): this may be a navigation page.")
    if "<pre" in str(cont) and art["counts"]["code_blocks"] == 0:
        conf -= 0.3
        warnings.append("Page contains <pre> blocks but no code was extracted.")
    art["confidence"] = round(max(0.0, min(1.0, conf)), 2)
    art["status"] = "ok" if art["confidence"] >= 0.6 and not warnings_blocking(warnings) else "warn"
    if art["confidence"] < 0.3:
        art["status"] = "failed"
        art["error"] = {"kind": "low_confidence", "message": f"Extraction confidence {art['confidence']} is too low.",
                        "hint": "Run 'inspect' on this page and tune --selector or the site profile."}
    return art


def warnings_blocking(w: list[str]) -> bool:
    return any(x.startswith(("High link density", "Page contains <pre>", "No paragraphs")) for x in w)


# ======================================================================================
# Discovery
# ======================================================================================
NON_ARTICLE_PATH = re.compile(
    r"(\.(png|jpe?g|gif|svg|webp|ico|pdf|zip|gz|css|js|xml|rss|mp4|mp3)$|/(tag|tags|category|categories|author|authors|"
    r"login|signin|signup|register|search|contact|about|privacy|terms|feed|cart|account|user|profile|jobs|courses?)(/|$))",
    re.I)


@dataclass
class Chapter:
    index: int
    url: str
    title: str = ""
    depth: int = 0
    source: str = "discovered"
    path: Path | None = None           # for imported local files
    section: str = ""                  # e.g. "SQL Basics" (becomes a Part in the PDF)
    section_desc: str = ""



JUNK_SECTION = re.compile(r"^(related (articles?|posts?|tutorials?)|recommended.*|trending.*|latest trend.*|comments?|"
                          r"you may also like|more from.*|read next|also read)$", re.I)


def page_subtitle(soup: BeautifulSoup) -> str:
    """First real sentence under the page title (e.g. 'SQL is the standard language ...')."""
    h1 = soup.find("h1")
    cands = h1.find_all_next("p", limit=8) if h1 else soup.find_all("p", limit=8)
    for p in cands:
        txt = collapse_ws(p.get_text(" "))
        if len(txt) < 40 or re.match(r"(last updated|updated|published)", txt, re.I):
            continue
        if len(txt) > 220:
            cut = txt[:220].rfind(". ")
            txt = txt[:cut + 1] if cut > 60 else txt[:217].rstrip() + "..."
        return txt
    return ""


def _section_candidate(a: Tag, page_h1: str, ids: set | None, cache: dict, scopes: list | None = None) -> tuple[str, str]:
    cur = a
    for _ in range(8):
        if ids:
            h = cur.find_previous(lambda t: isinstance(t, Tag) and id(t) in ids)
        else:
            h = cur.find_previous(["h2", "h3", "h4"])
        if h is None:
            return "", ""
        if scopes and not any(sc is h or any(par is sc for par in h.parents) for sc in scopes):
            return "", ""                      # heading outside the topic list's own area: not a section of it
        txt = collapse_ws(h.get_text(" "))
        inner = h.find("a")
        if (not txt or len(txt) > 120 or txt.lower() == page_h1.lower() or h in a.parents
                or (inner is not None and collapse_ws(inner.get_text(" ")) == txt)):
            cur = h
            continue
        if id(h) not in cache:
            desc = ""
            d = h.find_next(lambda t: isinstance(t, Tag) and t.name in ("p", "div", "span") and t is not h
                            and not t.find(["a", "ul", "ol", "h2", "h3", "h4", "table", "p", "div"])
                            and len(collapse_ws(t.get_text(" "))) > 25)
            if d is not None and any(x is d for x in a.find_all_previous(limit=4000)) and d not in a.parents:
                desc = collapse_ws(d.get_text(" "))
            cache[id(h)] = desc
        return txt, cache[id(h)]
    return "", ""


def assign_sections(anchors: list, soup: BeautifulSoup, info: dict, selector: str | None,
                    scopes: list | None = None) -> list[tuple[str, str]]:
    """Group links under the heading that precedes them. Tries h2-h4 first, then generic heading markup."""
    none = [("", "")] * len(anchors)
    if not anchors:
        return none
    page_h1 = info.get("page_title", "")
    attempts: list[tuple[str, set | None]] = []
    if selector:
        try:
            attempts.append((f"selector {selector}", {id(e) for e in soup.select(selector)}))
        except Exception:
            info["notes"].append(f"Bad --section-selector '{selector}' ignored.")
    attempts.append(("h2-h4 headings", None))
    for sel in ("[role=heading]", "[class*=eading]"):
        try:
            attempts.append((sel, {id(e) for e in soup.select(sel)}))
        except Exception:
            pass
    for label, ids in attempts:
        if ids is not None and not ids:
            continue
        cache: dict = {}
        res = [_section_candidate(a, page_h1, ids, cache, scopes) for a in anchors]
        names = [r[0] for r in res if r[0]]
        distinct = set(names)
        trusted = label.startswith("selector") or label == "h2-h4 headings"
        enough = len(names) >= max(3, len(anchors) // 2) and (len(distinct) >= 2 or label.startswith("selector"))
        if enough and (trusted or len(distinct) <= max(2, len(anchors) // 2)):   # generic markup: guard against nonsense
            info["sections_by"] = label
            return res
    return none


def sections_of(chapters: list) -> list[str]:
    out: list[str] = []
    for c in chapters:
        if c.section and c.section not in out:
            out.append(c.section)
    return out


def filter_sections(chapters: list, spec: str | None = None, skip: str | None = None) -> list:
    """Keep chapters in the chosen sections. spec: '1-5,8' or 'not 12,13' (section numbers as listed)."""
    secs = sections_of(chapters)
    if not secs:
        return chapters
    keep = set(secs)
    if spec and spec.strip():
        neg = bool(re.match(r"\s*not\b", spec, re.I))
        nums = parse_selection(re.sub(r"^\s*not\s*", "", spec, flags=re.I), len(secs))
        keep = {n for i, n in enumerate(secs, 1) if (i in nums) != neg}
    if skip:
        rx = re.compile(skip, re.I)
        keep = {n for n in keep if not rx.search(n)}
    return [c for c in chapters if not c.section or c.section in keep]


SIDEBAR_SEL = ("nav, aside, [class*=sidebar], [class*=Sidebar], [id*=sidebar], [id*=Sidebar], [class*=menu], [id*=menu], "
               "[class*=toc], [id*=toc], [class*=Toc], [class*=TOC], [class*=leftmenu], [class*=chapters], "
               "[class*=toctree]")


def count_links(anchors: list, base: str) -> int:
    keys = set()
    for a in anchors:
        u = abs_url(a.get("href"), base)
        if not u or not same_site(u, base) or NON_ARTICLE_PATH.search(urllib.parse.urlsplit(u).path):
            continue
        keys.add(dedupe_key(u))
    return len(keys)


def find_sidebar(soup: BeautifulSoup, base: str):
    """Find a tutorial's side menu (W3Schools / JavaTpoint / TutorialsPoint / docs style).
    Returns (element, links_share_the_page_directory, directory) or (None, False, '')."""
    here = urllib.parse.urlsplit(base)
    dirpath = here.path.rsplit("/", 1)[0] + "/"
    scored = []
    for el in soup.select(SIDEBAR_SEL):
        keys, same_dir = set(), 0
        for a in el.find_all("a", href=True):
            u = abs_url(a["href"], base)
            if not u or not same_site(u, base):
                continue
            p = urllib.parse.urlsplit(normalize_url(u))
            if NON_ARTICLE_PATH.search(p.path):
                continue
            k = dedupe_key(u)
            if k in keys:
                continue
            keys.add(k)
            if dirpath != "/" and p.path.startswith(dirpath):
                same_dir += 1
        if len(keys) >= 6:
            scored.append((el, len(keys), same_dir))
    if not scored:
        return None, False, ""
    coherent = [x for x in scored if x[2] >= 6 and x[2] >= 0.6 * x[1]]
    pool = coherent or [x for x in scored if x[1] <= 150]      # skip site-wide mega menus
    if not pool:
        return None, False, ""
    top = max(x[1] for x in pool)
    best = min((x for x in pool if x[1] >= 0.9 * top), key=lambda x: len(x[0].get_text()))
    return best[0], bool(coherent), dirpath


def same_site(a: str, b: str) -> bool:
    strip = lambda h: h[4:] if h.startswith("www.") else h
    return strip((urllib.parse.urlsplit(a).hostname or "").lower()) == strip((urllib.parse.urlsplit(b).hostname or "").lower())


def discover(f: Fetcher, index_url: str, s: Settings, link_selector: str | None = None,
             prefix: str | None = None) -> tuple[list[Chapter], dict]:
    res = f.get(index_url)
    soup = BeautifulSoup(res.body, "lxml")
    base = res.final_url
    bt = soup.find("base", href=True)
    if bt:
        base = urllib.parse.urljoin(base, bt["href"])
    profile = load_profile(urllib.parse.urlsplit(base).hostname or "", s)
    info = {"index_url": res.final_url, "confidence": "", "method": "", "notes": [],
            "page_title": page_title(soup, None), "page_subtitle": page_subtitle(soup)}
    scopes: list[Tag] = []
    anchors: list[Tag] = []
    if link_selector:
        for el in soup.select(link_selector):
            anchors += [el] if el.name == "a" else el.find_all("a", href=True)
        info.update(confidence="high", method=f"link-selector:{link_selector}")
    else:
        for sel in profile.get("toc", []):
            for el in soup.select(sel):
                scopes.append(el)
        if scopes:
            info.update(confidence="high", method="profile toc selectors")
        else:
            cont, m, _ = find_container(soup, profile, s.selector)
            if cont is not None:
                scopes = [cont]
                info.update(confidence="medium", method=f"content container ({m})")
        for sc in scopes:
            anchors += sc.find_all("a", href=True)
        if not info["method"].startswith("profile") and count_links(anchors, base) < 6:
            side, coherent, dirpath = find_sidebar(soup, base)
            if side is not None:
                scopes, anchors = [side], side.find_all("a", href=True)
                info.update(confidence="medium", method="sidebar menu")
                if coherent and not prefix:
                    prefix = dirpath
        if len(anchors) < 3 and soup.body is not None:
            body = copy.copy(soup.body)
            for e in body.find_all(["header", "footer", "nav"]):
                e.decompose()
            anchors = body.find_all("a", href=True)
            info.update(confidence="low", method="whole page (header/footer/nav removed)")
            info["notes"].append("Few links in the content area; fell back to the whole page. Review the list carefully.")
    li_anchors = [a for a in anchors if a.find_parent("li")]
    if len(li_anchors) >= 5:
        anchors = li_anchors
    seen: dict[str, Chapter] = {}
    self_key = dedupe_key(res.final_url)
    include_self = bool(link_selector) or info["method"].startswith(("profile", "sidebar"))
    min_depth = None
    raw: list[tuple] = []
    for a in anchors:
        u = abs_url(a.get("href"), base)
        if not u or not same_site(u, base):
            continue
        try:
            u = normalize_url(u)
        except Exception:
            continue
        p = urllib.parse.urlsplit(u)
        if NON_ARTICLE_PATH.search(p.path) or (not include_self and dedupe_key(u) == self_key):
            continue
        if prefix and not p.path.startswith(prefix):
            continue
        title = collapse_ws(a.get_text(" "))
        if not title or len(title) > 160:
            continue
        depth = len(a.find_parents(["ul", "ol"]))
        min_depth = depth if min_depth is None else min(min_depth, depth)
        raw.append((u, title, depth, a))
    secs = assign_sections([r[3] for r in raw], soup, info, s.section_selector, scopes)
    dup = 0
    for (u, title, depth, _a), (sec, sdesc) in zip(raw, secs):
        k = dedupe_key(u)
        if k in seen:
            dup += 1
            continue
        seen[k] = Chapter(0, u, title, max(0, depth - (min_depth or 0)), section=sec, section_desc=sdesc)
    chapters = list(seen.values())
    dropped = sorted({c.section for c in chapters if c.section and JUNK_SECTION.match(c.section)})
    if dropped:
        chapters = [c for c in chapters if c.section not in dropped]
        info["notes"].append("Left out non-tutorial section(s): " + ", ".join(dropped) + ".")
    if len(chapters) > s.max_pages:
        info["notes"].append(f"Found {len(chapters)} links; truncated to --max-pages {s.max_pages}.")
        chapters = chapters[:s.max_pages]
    for i, c in enumerate(chapters, 1):
        c.index = i
    if dup:
        info["notes"].append(f"{dup} duplicate link(s) removed.")
    if info["method"].startswith(("sidebar", "profile")) and self_key in seen:      # a chapter page was pasted
        first = next((c for c in chapters if c.section), None)
        if first:
            info["page_title"], info["page_subtitle"] = first.section, first.section_desc
        else:
            info["page_subtitle"] = ""
    return chapters, info


def write_chapters_file(path: Path, chapters: list[Chapter], info: dict) -> None:
    lines = [
        f"# TutorialStudyForge (TSF) topic list  (generated {datetime.now().isoformat(timespec='seconds')})",
        f"# Source index: {info.get('index_url', '')}",
        f"# Title: {info.get('page_title', '')}",
        f"# Subtitle: {info.get('page_subtitle', '')}",
        f"# Discovery: {info.get('method', '')} (confidence: {info.get('confidence', '')})",
        "# EDIT ME: delete a line to leave a topic out; move lines to reorder them.",
        "# Lines starting with '==' are section headings (they become Parts in the PDF).",
        "# Format:  URL  # title",
    ] + [f"# NOTE: {n}" for n in info.get("notes", [])]
    has_sec = any(c.section for c in chapters)
    cur = None
    for c in chapters:
        if has_sec and c.section != cur:
            cur = c.section
            if cur:
                lines += ["", f"== {cur}" + (f"  # {c.section_desc}" if c.section_desc else "")]
        indent = "" if has_sec else "  " * c.depth
        lines.append(f"{indent}{c.url}  # {c.title}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def read_chapters_file(path: Path) -> list[Chapter]:
    """Parse a topic list. Metadata (# Title:, # Subtitle:) is left in read_chapters_file.meta."""
    chapters: list[Chapter] = []
    meta: dict = {}
    sec, desc = "", ""
    for raw in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"#\s*(Title|Subtitle):\s*(.*)$", raw)
        if m:
            meta[m.group(1).lower()] = m.group(2).strip()
        if not raw.strip() or raw.lstrip().startswith("#"):
            continue
        body = raw.strip()
        if body.startswith("=="):
            name, _, d = body[2:].strip().partition("  # ")
            sec, desc = name.strip(), d.strip()
            continue
        depth = (len(raw) - len(raw.lstrip(" "))) // 2
        url, title = (body.split("  # ", 1) + [""])[:2] if "  # " in body else (body, "")
        chapters.append(Chapter(len(chapters) + 1, url.strip(), title.strip(), 0 if sec else depth, source="file",
                                section=sec, section_desc=desc))
    read_chapters_file.meta = meta
    return chapters


read_chapters_file.meta = {}


def parse_selection(spec: str, n: int) -> set[int]:
    out: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise ValueError(f"Bad --select item '{part}' (use e.g. 1-5,8)")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        out.update(range(a, b + 1))
    bad = [i for i in out if i < 1 or i > n]
    if bad:
        raise ValueError(f"--select index out of range: {sorted(bad)} (1..{n})")
    return out


# ======================================================================================
# Images
# ======================================================================================
MAGIC = [(b"\x89PNG\r\n\x1a\n", "png"), (b"\xff\xd8\xff", "jpg"), (b"GIF87a", "gif"), (b"GIF89a", "gif")]


def sniff_image(data: bytes) -> str | None:
    for m, ext in MAGIC:
        if data.startswith(m):
            return ext
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    return None


def attach_images(art: dict, fetcher: Fetcher | None, assets: Path, s: Settings, local_dir: Path | None,
                  budget: dict) -> None:
    def rec(bs):
        for b in bs:
            if b["type"] == "image":
                handle(b)
            elif b["type"] == "list":
                for it in b["items"]:
                    rec(it)
            elif b["type"] == "quote":
                rec(b["blocks"])

    count = 0

    def handle(b: dict) -> None:
        nonlocal count
        if not s.images:
            return
        if count >= s.max_images_per_article or budget["left"] <= 0:
            art["warnings"].append("Image limit reached; some images were not downloaded.")
            return
        count += 1
        budget["left"] -= 1
        try:
            if local_dir is not None and not re.match(r"^[a-z]+:", b["src"], re.I) and not b["src"].startswith("//"):
                p = (local_dir / urllib.parse.unquote(b["src"].split("#")[0].split("?")[0])).resolve()
                if local_dir.resolve() not in p.parents:
                    raise FetchError("ssrf", "Image path escapes the import folder.")
                data = p.read_bytes()
                if len(data) > s.max_image_bytes:
                    raise FetchError("too_large", "Image too large.")
            elif b.get("url") and fetcher is not None:
                data = fetcher.get(b["url"], kind="image").body
            else:
                raise FetchError("url", "No usable image URL.")
            ext = sniff_image(data)
            if not ext:
                raise FetchError("bad_type", "Not a supported image type (png/jpg/gif/webp).")
            if len(data) < 800:
                b["asset"] = None
                b["skipped"] = True
                return
            name = f"{hashlib.sha256(data).hexdigest()[:16]}.{ext}"
            assets.mkdir(parents=True, exist_ok=True)
            (assets / name).write_bytes(data)
            b["asset"] = name
        except (FetchError, OSError) as e:
            msg = e.message if isinstance(e, FetchError) else e.__class__.__name__
            art["warnings"].append(f"Image unavailable ({msg}): {b['src'][:80]}")
            b["asset"] = None

    rec(art["blocks"])


# ======================================================================================
# Rendering: blocks -> HTML / Markdown
# ======================================================================================
CONT_MARK = "\u21b3"


def wrap_code(text: str, cols: int) -> tuple[str, int]:
    """Wrap overlong code lines with a visible continuation marker. Returns (text, wrapped_line_count)."""
    out: list[str] = []
    wrapped = 0
    for line in text.expandtabs(4).split("\n"):
        if len(line) <= cols:
            out.append(line)
            continue
        indent = re.match(r"\s*", line).group(0)
        cont_prefix = indent + CONT_MARK + " "
        rest = line
        first = True
        while len(rest) > (cols if first else cols):
            width = cols if first else cols
            window = rest[:width]
            cut = max(window.rfind(" "), window.rfind(","), window.rfind(";"), window.rfind(")"))
            if cut < width * 0.4:
                cut = width - 1
            out.append(rest[:cut + 1].rstrip() if first else rest[:cut + 1].rstrip())
            rest = cont_prefix + rest[cut + 1:].lstrip()
            first = False
            wrapped += 1
            if len(cont_prefix) >= cols - 10:   # pathological indent; stop wrapping to avoid infinite loop
                break
        out.append(rest)
    return "\n".join(out), wrapped


class HtmlRenderer:
    def __init__(self, s: Settings, assets_rel: str = "assets"):
        self.s = s
        self.assets_rel = assets_rel
        self.wrapped_lines = 0
        self.h_counter = 0
        self.headings: list[tuple[int, str, str]] = []   # (level, id, plain text)

    def blocks(self, bs: list[dict], chapter_no: int, depth: int = 0) -> str:
        return "".join(self.block(b, chapter_no, depth) for b in bs)

    def block(self, b: dict, ch: int, depth: int) -> str:
        t = b["type"]
        if t == "paragraph":
            return f"<p>{b['html']}</p>\n"
        if t == "heading":
            self.h_counter += 1
            hid = f"h-{ch}-{self.h_counter}"
            plain = re.sub(r"<[^>]+>", "", htmllib.unescape(b["html"]))
            self.headings.append((b["level"], hid, plain))
            return f'<h{b["level"]} id="{hid}">{b["html"]}</h{b["level"]}>\n'
        if t == "code":
            cols = max(40, self.s.code_cols - 6 * depth)
            txt, n = wrap_code(b["text"], cols)
            self.wrapped_lines += n
            lang = f'<div class="code-lang">{esc(b["language"])}</div>' if b["language"] else ""
            return f'<div class="code-wrap">{lang}<pre class="code"><code>{esc(txt)}</code></pre></div>\n'
        if t == "list":
            tag = "ol" if b["ordered"] else "ul"
            start = f' start="{b["start"]}"' if b["ordered"] and b["start"] != 1 else ""
            items = "".join(f"<li>{self.blocks(it, ch, depth + 1)}</li>\n" for it in b["items"])
            return f"<{tag}{start}>\n{items}</{tag}>\n"
        if t == "table":
            return self.table(b)
        if t == "image":
            if b.get("asset"):
                cap = f"<figcaption>{esc(b['caption'])}</figcaption>" if b["caption"] else ""
                return (f'<figure><img src="{self.assets_rel}/{b["asset"]}" alt="{htmllib.escape(b["alt"], quote=True)}">'
                        f"{cap}</figure>\n")
            if b.get("skipped"):
                return ""
            label = b["caption"] or b["alt"] or "image"
            return f'<p class="missing">[Image not available: {esc(label)}]</p>\n'
        if t == "quote":
            kind = f' class="callout callout-{esc(b["kind"])}"' if b["kind"] else ""
            label = f'<div class="callout-label">{esc(b["kind"].title())}</div>' if b["kind"] else ""
            return f"<blockquote{kind}>{label}{self.blocks(b['blocks'], ch, depth + 1)}</blockquote>\n"
        if t == "equation":
            return f'<pre class="equation">{esc(b["tex"])}</pre>\n'
        if t == "hr":
            return "<hr>\n"
        return ""

    def table(self, b: dict) -> str:
        cap = f"<caption>{esc(b['caption'])}</caption>" if b["caption"] else ""
        head = [r for r in b["rows"] if r["head"]]
        body = [r for r in b["rows"] if not r["head"]]

        def row(r):
            cells = ""
            for c in r["cells"]:
                tag = "th" if c["th"] else "td"
                sp = (f' colspan="{c["colspan"]}"' if c["colspan"] > 1 else "") + \
                     (f' rowspan="{c["rowspan"]}"' if c["rowspan"] > 1 else "")
                cells += f"<{tag}{sp}>{c['html']}</{tag}>"
            return f"<tr>{cells}</tr>"
        th = f"<thead>{''.join(row(r) for r in head)}</thead>" if head else ""
        return f"<table>{cap}{th}<tbody>{''.join(row(r) for r in body)}</tbody></table>\n"


def css_for(s: Settings, title: str) -> str:
    page = f"{s.page_size} {'landscape' if s.landscape else 'portrait'}"
    w_mm = {"A4": 210.0, "LETTER": 215.9}.get(s.page_size.upper(), 210.0)
    h_mm = {"A4": 297.0, "LETTER": 279.4}.get(s.page_size.upper(), 297.0)
    if s.landscape:
        w_mm = h_mm
    content_pt = (w_mm - 2 * s.margin_mm) * 72 / 25.4
    code_pt = max(6.0, min(s.font_pt - 1.0, content_pt / (s.code_cols * 0.62)))
    m = s.margin_mm
    return f"""
@page {{ size: {page}; margin: {m + 2}mm {m}mm {m + 4}mm {m}mm;
  @top-left {{ content: string(chaptername); font: 8pt "DejaVu Sans", sans-serif; color: #555; }}
  @top-right {{ content: "{title.replace('"', "'")[:60]}"; font: 8pt "DejaVu Sans", sans-serif; color: #555; }}
  @bottom-center {{ content: counter(page); font: 9pt "DejaVu Sans", sans-serif; color: #333; }} }}
@page cover {{ @top-left {{ content: none }} @top-right {{ content: none }} @bottom-center {{ content: none }} }}
html {{ font: {s.font_pt}pt/1.45 "DejaVu Serif", "Noto Serif", Georgia, serif; color: #111; }}
body {{ hyphens: manual; }}
.cover {{ page: cover; break-after: page; text-align: center; padding-top: 70mm; }}
.cover h1 {{ font: 700 28pt/1.2 "DejaVu Sans", sans-serif; bookmark-level: none; margin: 0 0 8mm; }}
.cover .subtitle {{ font: 14pt "DejaVu Sans", sans-serif; color: #444; }}
.cover .meta {{ margin-top: 25mm; font: 10pt "DejaVu Sans", sans-serif; color: #444; }}
.cover .notice {{ margin: 18mm auto 0; max-width: 120mm; font: 8.5pt/1.4 "DejaVu Sans", sans-serif; color: #555; }}
.cover .partial {{ margin-top: 10mm; font: 700 12pt "DejaVu Sans", sans-serif; color: #900; }}
h1, h2, h3, h4, h5, h6 {{ font-family: "DejaVu Sans", sans-serif; break-after: avoid; line-height: 1.25; }}
h1.chapter-title {{ font-size: 20pt; break-before: page; bookmark-level: 1; string-set: chaptername content(); margin: 0 0 3mm; }}
.src {{ font: 8pt "DejaVu Sans", sans-serif; color: #666; margin: 0 0 6mm; word-break: break-all; }}
h2 {{ font-size: 15pt; bookmark-level: 2; margin: 7mm 0 2mm; }} h3 {{ font-size: 12.5pt; bookmark-level: 3; margin: 5mm 0 2mm; }}
h4 {{ font-size: 11pt; bookmark-level: 4; }} h5, h6 {{ font-size: 10.5pt; bookmark-level: none; }}
p {{ margin: 0 0 3mm; orphans: 3; widows: 3; }}
a {{ color: #0b4a8f; text-decoration: none; }}
li > p {{ margin: 0 0 1.5mm; }} ul, ol {{ margin: 0 0 3mm; padding-left: 7mm; }}
code {{ font-family: "DejaVu Sans Mono", "Liberation Mono", monospace; font-size: 0.9em; background: #f1f1f1; padding: 0 1px; }}
.code-wrap {{ margin: 0 0 4mm; }}
.code-lang {{ font: 700 7pt "DejaVu Sans", sans-serif; color: #444; text-transform: uppercase; letter-spacing: .05em; margin-bottom: 1px; }}
pre.code {{ margin: 0; padding: 2.5mm 3mm; background: #f6f6f6; border: 0.6pt solid #bbb; border-left: 2pt solid #444;
  white-space: pre; font: {code_pt:.2f}pt/1.35 "DejaVu Sans Mono", "Liberation Mono", monospace; orphans: 3; widows: 3; }}
pre.code code {{ background: none; padding: 0; font-size: 1em; }}
pre.equation {{ font: 9.5pt "DejaVu Sans Mono", monospace; white-space: pre-wrap; margin: 0 0 4mm; }}
table {{ border-collapse: collapse; width: 100%; margin: 0 0 4mm; font-size: 0.92em; }}
caption {{ font-weight: 700; text-align: left; margin-bottom: 1mm; caption-side: top; }}
th, td {{ border: 0.6pt solid #999; padding: 1.2mm 2mm; vertical-align: top; text-align: left; overflow-wrap: anywhere; }}
th {{ background: #ececec; font-family: "DejaVu Sans", sans-serif; }} thead {{ display: table-header-group; }} tr {{ break-inside: avoid; }}
figure {{ margin: 0 0 4mm; text-align: center; break-inside: avoid; }} figure img {{ max-width: 100%; max-height: 120mm; }}
figcaption {{ font: italic 8.5pt "DejaVu Sans", sans-serif; color: #444; margin-top: 1mm; }}
blockquote {{ margin: 0 0 4mm; padding: 1mm 4mm; border-left: 2pt solid #999; color: #222; }}
.callout-label {{ font: 700 8pt "DejaVu Sans", sans-serif; text-transform: uppercase; color: #444; }}
.missing {{ font-style: italic; color: #666; }}
.failed {{ border: 1pt dashed #900; padding: 4mm; color: #600; font-family: "DejaVu Sans", sans-serif; }}
.toc h1 {{ font-size: 20pt; bookmark-level: 1; break-before: page; }}
.toc ul {{ list-style: none; padding-left: 0; }} .toc li {{ margin: 1mm 0; }} .toc li.d1 {{ padding-left: 6mm; font-size: .92em; }}
.toc li.d2 {{ padding-left: 12mm; font-size: .88em; }}
.toc a {{ color: #111; }} .toc a::after {{ content: leader('.') target-counter(attr(href), page); }}
h1.appendix {{ font-size: 20pt; break-before: page; bookmark-level: 1; string-set: chaptername content(); }}
table.report {{ font: 8pt "DejaVu Sans", sans-serif; }} .st-failed {{ color: #900; font-weight: 700; }} .st-warn {{ color: #8a5a00; font-weight: 700; }}
"""


PARTS_CSS = """
.part { break-before: page; break-after: page; padding-top: 28mm; }
.part-label { font: 700 9pt "DejaVu Sans", sans-serif; letter-spacing: .14em; text-transform: uppercase; color: #666; margin-bottom: 2mm; }
h1.part-title { font-size: 26pt; bookmark-level: 1; string-set: chaptername content(); margin: 0 0 5mm; border-bottom: 2pt solid #444; padding-bottom: 3mm; }
.part-desc { font: italic 11.5pt/1.5 "DejaVu Serif", serif; color: #333; margin: 0 0 8mm; }
ol.part-list { list-style: none; padding: 0; margin: 0; font: 10.5pt "DejaVu Sans", sans-serif; }
ol.part-list li { margin: 1.6mm 0; } ol.part-list a { color: #111; }
ol.part-list a::after { content: leader('.') target-counter(attr(href), page); }
body.parts h1.chapter-title { bookmark-level: 2; font-size: 18pt; }
body.parts h2 { bookmark-level: 3; } body.parts h3 { bookmark-level: 4; } body.parts h4 { bookmark-level: 5; }
.toc li.toc-part { margin-top: 3mm; font-weight: 700; font-family: "DejaVu Sans", sans-serif; }
.toc li.d3 { padding-left: 18mm; font-size: .85em; }
"""


def build_html(articles: list[dict], s: Settings, meta: dict, status: str, report_rows: list[dict],
               assets_rel: str = "assets") -> tuple[str, HtmlRenderer]:
    r = HtmlRenderer(s, assets_rel)
    title = esc(meta["title"])
    has_parts = any(a.get("section") for a in articles)
    toc_depth = s.toc_depth if s.toc_depth is not None else (1 if has_parts else 2)
    parts: list[str] = []
    if s.cover:
        sub = f'<div class="subtitle">{esc(s.subtitle)}</div>' if s.subtitle else ""
        by = f"<div>{esc(s.author)}</div>" if s.author else ""
        partial = '<div class="partial">PARTIAL DOCUMENT: some chapters could not be retrieved (see build report)</div>' \
            if status == "Partial" else ""
        n_ok = len([a for a in articles if a["status"] in ("ok", "warn")])
        n_parts = len({a["section"] for a in articles if a.get("section")})
        counts = f"{n_ok} chapters" + (f" in {n_parts} parts" if has_parts else "")
        parts.append(
            f'<section class="cover"><h1>{title}</h1>{sub}<div class="meta">{by}<div>Compiled {meta["date"]}</div>'
            f'<div>{counts}</div></div>{partial}'
            f'<div class="notice">Compiled for personal study from the sources listed at the end of this document. '
            f"Content belongs to its original authors; do not redistribute. Faithful compilation: text was not "
            f"rewritten or summarised.</div></section>")
    body: list[str] = []
    toc_items: list[str] = []
    cur_sec, part_no = None, 0
    for a in articles:
        n = a["index"]
        cid = f"ch-{n}"
        t = esc(a["title"] or f"Chapter {n}")
        sec = a.get("section") or ""
        if has_parts and sec and sec != cur_sec:
            cur_sec, part_no = sec, part_no + 1
            pid = f"part-{part_no}"
            members = [x for x in articles if (x.get("section") or "") == sec]
            lis = "".join(f'<li><a href="#ch-{x["index"]}">{esc(x["title"] or "Chapter %d" % x["index"])}</a></li>'
                          for x in members)
            desc = f'<p class="part-desc">{esc(a["section_desc"])}</p>' if a.get("section_desc") else ""
            body.append(f'<section class="part"><div class="part-label">Part {part_no}</div>'
                        f'<h1 class="part-title" id="{pid}">{esc(sec)}</h1>{desc}<ol class="part-list">{lis}</ol></section>')
            toc_items.append(f'<li class="d0 toc-part"><a href="#{pid}">{esc(sec)}</a></li>')
        base = 1 if (has_parts and sec) else 0
        start_h = len(r.headings)
        if a["status"] in ("ok", "warn"):
            src = f'<p class="src">Source: <a href="{htmllib.escape(a["url"], quote=True)}">{esc(a["url"])}</a></p>' \
                if s.source_urls and a["url"].startswith("http") else ""
            body.append(f'<section class="chapter"><h1 class="chapter-title" id="{cid}">{t}</h1>{src}'
                        f'{r.blocks(a["blocks"], n)}</section>')
        else:
            err = a.get("error") or {}
            body.append(f'<section class="chapter"><h1 class="chapter-title" id="{cid}">{t}</h1>'
                        f'<div class="failed"><strong>This chapter is missing from this document.</strong><br>'
                        f'Reason: {esc(err.get("message", a["status"]))}<br>Source: {esc(a["url"])}</div></section>')
        toc_items.append(f'<li class="d{base}"><a href="#{cid}">{t}</a></li>')
        if toc_depth >= 2:
            for lvl, hid, txt in r.headings[start_h:]:
                if lvl <= toc_depth + 1:
                    toc_items.append(f'<li class="d{min(lvl - 1 + base, 3)}"><a href="#{hid}">{esc(txt)}</a></li>')
    if s.toc:
        parts.append(f'<section class="toc"><h1>Contents</h1><ul>{"".join(toc_items)}</ul></section>')
    parts += body
    if s.source_urls:
        lis = "".join(f'<li>{esc(a["title"] or a["url"])}: <a href="{htmllib.escape(a["url"], quote=True)}">{esc(a["url"])}</a></li>'
                      for a in articles if a["url"].startswith("http"))
        if lis:
            parts.append(f'<section><h1 class="appendix">Sources</h1><p>Original locations of the compiled chapters.</p><ul>{lis}</ul></section>')
    if s.report_page:
        rows = "".join(
            f'<tr><td>{x["index"]}</td><td>{esc(x["title"][:70])}</td><td class="st-{x["status"]}">{x["status"]}</td>'
            f'<td>{x["confidence"]}</td><td>{x["counts"].get("code_blocks", 0)}</td><td>{x["counts"].get("tables", 0)}</td>'
            f'<td>{esc("; ".join(x["notes"])[:160])}</td></tr>' for x in report_rows)
        parts.append(f'<section><h1 class="appendix">Build report</h1><p>Status: <strong>{esc(status)}</strong>. '
                     f'Generated by {APP_NAME} (TSF) {__version__}. Code lines wrapped: {r.wrapped_lines}.</p>'
                     f'<table class="report"><thead><tr><th>#</th><th>Chapter</th><th>Status</th><th>Conf.</th><th>Code</th>'
                     f'<th>Tables</th><th>Notes</th></tr></thead><tbody>{rows}</tbody></table></section>')
    css = css_for(s, meta["title"]) + PARTS_CSS + \
        ("" if s.chapter_breaks else "h1.chapter-title { break-before: auto; margin-top: 8mm; }")
    doc = (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{title}</title>'
           f'<meta name="author" content="{htmllib.escape(s.author or "", quote=True)}">'
           f'<meta name="description" content="{htmllib.escape((s.subtitle or "Compiled study guide") + " (" + status + ")", quote=True)}">'
           f'<meta name="generator" content="{APP_NAME} (TSF) {__version__}">'
           f'<style>{css}</style></head><body{" class=parts" if has_parts else ""}>{"".join(parts)}</body></html>')
    return doc, r


# ---- Markdown ----
def md_inline(html_str: str) -> str:
    soup = BeautifulSoup(html_str, "html.parser")

    def conv(node) -> str:
        if isinstance(node, NavigableString):
            return re.sub(r"([\\`*_\[\]<>])", r"\\\1", str(node))
        n = node.name
        inner = "".join(conv(c) for c in node.children)
        if n == "strong":
            return f"**{inner}**"
        if n == "em":
            return f"*{inner}*"
        if n == "code":
            t = node.get_text()
            fence = "`" * (max([len(m) for m in re.findall(r"`+", t)] + [0]) + 1)
            pad = " " if t.startswith("`") or t.endswith("`") else ""
            return f"{fence}{pad}{t}{pad}{fence}"
        if n == "a":
            return f"[{inner}]({node.get('href', '')})"
        if n == "br":
            return "  \n"
        if n in ("sub", "sup"):
            return f"<{n}>{inner}</{n}>"
        return inner
    return "".join(conv(c) for c in soup.children).strip()


def md_blocks(bs: list[dict], indent: str = "", hshift: int = 0) -> list[str]:
    out: list[str] = []
    for b in bs:
        t = b["type"]
        if t == "paragraph":
            out += [indent + md_inline(b["html"]), ""]
        elif t == "heading":
            out += [f"{'#' * min(6, b['level'] + hshift)} {md_inline(b['html'])}", ""]
        elif t == "code":
            fence = "`" * max(3, max([len(m) for m in re.findall(r"`+", b["text"])] + [0]) + 1)
            out += [indent + fence + b["language"]] + [indent + ln for ln in b["text"].split("\n")] + [indent + fence, ""]
        elif t == "list":
            n = b["start"]
            for it in b["items"]:
                marker = f"{n}." if b["ordered"] else "-"
                n += 1
                sub = md_blocks(it, indent + " " * (len(marker) + 1), hshift)
                while sub and sub[-1] == "":
                    sub.pop()
                sub = [ln for k, ln in enumerate(sub) if not (ln == "" and k + 1 < len(sub)
                       and re.match(r"\s*(?:[-*]|\d+\.)\s", sub[k + 1]) and k > 0 and not sub[k - 1].lstrip().startswith("```"))]
                if sub:
                    sub[0] = indent + marker + " " + sub[0].lstrip()
                    out += sub
            out.append("")
        elif t == "table":
            simple = all(c["colspan"] == 1 and c["rowspan"] == 1 for r in b["rows"] for c in r["cells"])
            if simple:
                width = max(len(r["cells"]) for r in b["rows"])
                rows = [[md_inline(c["html"]).replace("|", "\\|").replace("\n", " ") for c in r["cells"]] +
                        [""] * (width - len(r["cells"])) for r in b["rows"]]
                if b["caption"]:
                    out += [f"**{b['caption']}**", ""]
                out += ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]] + [""]
            else:
                out += ["<!-- table with merged cells: kept as HTML -->", HtmlRenderer(Settings()).table(b), ""]
        elif t == "image":
            alt = (b["alt"] or b["caption"] or "image").replace("]", "")
            out += [f"![{alt}]({'assets/' + b['asset'] if b.get('asset') else (b.get('url') or b['src'])})"
                    + (f"\n*{b['caption']}*" if b["caption"] else ""), ""]
        elif t == "quote":
            inner = md_blocks(b["blocks"], hshift=hshift)
            while inner and inner[-1] == "":
                inner.pop()
            label = [f"**{b['kind'].title()}**", ""] if b["kind"] else []
            out += [indent + "> " + ln if ln else indent + ">" for ln in label + inner] + [""]
        elif t == "equation":
            out += ["```latex", b["tex"], "```", ""]
        elif t == "hr":
            out += ["---", ""]
    return out


def build_markdown(articles: list[dict], meta: dict, s: Settings, status: str) -> str:
    has_parts = any(a.get("section") for a in articles)
    lines = [f"# {meta['title']}", ""]
    if s.subtitle:
        lines += [f"*{s.subtitle}*", ""]
    lines += [f"Compiled {meta['date']} by {APP_NAME} (TSF) {__version__}. Status: **{status}**. "
              f"For personal study; content belongs to its original authors.", ""]
    cur = None
    for a in articles:
        sec = a.get("section") or ""
        in_part = has_parts and bool(sec)
        if in_part and sec != cur:
            cur = sec
            lines += [f"## {sec}", ""]
            if a.get("section_desc"):
                lines += [f"*{a['section_desc']}*", ""]
        lines += [f"{'###' if in_part else '##'} {a['title'] or 'Chapter ' + str(a['index'])}", ""]
        if a["url"].startswith("http"):
            lines += [f"<!-- source: {a['url']} -->", f"Source: <{a['url']}>", ""]
        if a["status"] in ("ok", "warn"):
            lines += md_blocks(a["blocks"], hshift=2 if in_part else 1)
        else:
            lines += [f"> **Missing chapter:** {(a.get('error') or {}).get('message', a['status'])}", ""]
    return "\n".join(lines).rstrip() + "\n"


# ======================================================================================
# PDF generation and verification
# ======================================================================================
def render_pdf(html_doc: str, out_pdf: Path, base_dir: Path) -> None:
    try:
        import weasyprint
    except ImportError:
        say("Missing dependency 'weasyprint'. Install with: pip install weasyprint  (needs Pango; see README)")
        sys.exit(1)
    except OSError as e:
        say(f"WeasyPrint could not load its system libraries (Pango): {e}\nSee https://doc.courtbouillon.org/weasyprint/stable/first_steps.html")
        sys.exit(1)
    kwargs: dict = {}
    if hasattr(weasyprint, "URLFetcher"):                          # WeasyPrint >= 68
        kwargs["url_fetcher"] = weasyprint.URLFetcher(allowed_protocols=["file", "data"])
    else:                                                          # older versions: function fetcher
        from weasyprint.urls import default_url_fetcher

        def safe_fetcher(url, *a, **k):
            if not url.startswith(("file:", "data:")):
                raise ValueError("External resources are blocked.")
            return default_url_fetcher(url, *a, **k)
        kwargs["url_fetcher"] = safe_fetcher
    weasyprint.HTML(string=html_doc, base_url=str(base_dir) + "/", **kwargs).write_pdf(str(out_pdf))


def _nospace(t: str) -> str:
    return re.sub(r"\s+", "", t.replace("\u00ad", "")).lower()


def verify_pdf(pdf: Path, articles: list[dict], meta: dict, s: Settings) -> list[dict]:
    """Independent checks on the finished PDF. Each check: {name, ok, detail}."""
    checks: list[dict] = []

    def add(name, ok, detail=""):
        checks.append({"name": name, "ok": bool(ok), "detail": detail})
    try:
        from pypdf import PdfReader
    except ImportError:
        add("pypdf available", False, "Install pypdf to enable PDF verification: pip install pypdf")
        return checks
    if not pdf.exists() or pdf.stat().st_size == 0:
        add("PDF exists and is non-empty", False, "file missing or empty")
        return checks
    add("PDF exists and is non-empty", True, f"{pdf.stat().st_size} bytes")
    try:
        rd = PdfReader(str(pdf))
        pages = len(rd.pages)
    except Exception as e:
        add("PDF can be parsed", False, str(e))
        return checks
    add("PDF can be parsed", pages > 0, f"{pages} pages")
    title = (rd.metadata.title if rd.metadata else None) or ""
    add("Metadata title set", title == meta["title"], f"title={title!r}")
    texts = [(p.extract_text() or "") for p in rd.pages]
    full = _nospace("".join(texts))
    add("Text is extractable (searchable)", len(full) > 200, f"{len(full)} characters")

    def flat(o):
        for x in o:
            if isinstance(x, list):
                yield from flat(x)
            else:
                yield x
    try:
        marks = [_nospace(getattr(x, "title", "") or "") for x in flat(rd.outline)]
    except Exception:
        marks = []
    included = [a for a in articles if a["status"] in ("ok", "warn")]
    missing_marks = [a["title"] for a in included if _nospace(a["title"]) not in marks]
    add("Bookmarks exist for every chapter", bool(marks) and not missing_marks,
        f"{len(marks)} bookmarks" + (f"; missing: {missing_marks[:3]}" if missing_marks else ""))
    miss_text, miss_code = [], []
    for a in included:
        txt = plain_text(a["blocks"])
        first_par = next((b for b in a["blocks"] if b["type"] == "paragraph"), None)
        if first_par:
            words = re.sub(r"<[^>]+>", "", htmllib.unescape(first_par["html"])).split()[:7]
            if _nospace(" ".join(words)) not in full:
                miss_text.append(a["title"])
        codes = [ln.strip() for b in a["blocks"] if b["type"] == "code" for ln in b["text"].split("\n")
                 if 10 <= len(ln.strip()) <= min(60, s.code_cols - 12)]
        if codes and _nospace(codes[0]) not in full:
            miss_code.append(a["title"])
        if a["counts"].get("code_blocks") and not codes:
            pass
    add("Opening text of every chapter found in PDF", not miss_text, f"missing: {miss_text[:3]}" if miss_text else "")
    add("Sample code line of every chapter found in PDF", not miss_code, f"missing: {miss_code[:3]}" if miss_code else "")
    if s.toc:
        add("Table of contents present", _nospace("Contents") in full)
    return checks


# ======================================================================================
# Pipeline
# ======================================================================================
def process_chapters(chapters: list[Chapter], fetcher: Fetcher | None, s: Settings, assets: Path,
                     local_mode: bool, base_url: str | None) -> list[dict]:
    articles: list[dict] = []
    seen_hash: dict[str, int] = {}
    seen_url: dict[str, int] = {}
    budget = {"left": 300}
    total = len(chapters)
    for ch in chapters:
        say(f"[{ch.index}/{total}] {ch.title or ch.url or ch.path}")
        art: dict = {"schema": SCHEMA_VERSION, "index": ch.index, "url": ch.url, "canonical_url": None,
                     "final_url": None, "title": ch.title, "section": ch.section,
                     "section_desc": ch.section_desc, "source_name": "user-supplied" if local_mode else
                     (urllib.parse.urlsplit(ch.url).hostname or ""), "depth": ch.depth, "status": "failed",
                     "confidence": 0.0, "method": "", "fetched_at": None, "from_cache": False, "content_hash": "",
                     "words": 0, "counts": {}, "warnings": [], "error": None, "blocks": []}
        local_dir = None
        try:
            if local_mode:
                data = ch.path.read_bytes()
                local_dir = ch.path.parent
                final_url, fetched_at = ch.url, datetime.fromtimestamp(ch.path.stat().st_mtime, timezone.utc).isoformat(timespec="seconds")
                base = base_url
            else:
                assert fetcher is not None
                k = dedupe_key(ch.url)
                if k in seen_url:
                    art.update(status="duplicate", error={"kind": "duplicate", "message": f"Same URL as chapter {seen_url[k]}.", "hint": ""})
                    articles.append(art)
                    continue
                res = fetcher.get(ch.url)
                data, final_url, fetched_at, base = res.body, res.final_url, res.fetched_at, res.final_url
                art["from_cache"] = res.from_cache
                seen_url[k] = ch.index
                seen_url[dedupe_key(final_url)] = ch.index
            art["final_url"], art["fetched_at"] = final_url, fetched_at
            ex = extract_article(data, final_url, base, s, profile_host=urllib.parse.urlsplit(final_url).hostname or "")
            for key in ("blocks", "warnings", "method", "confidence", "status", "error", "canonical_url", "words",
                        "content_hash", "counts"):
                art[key] = ex[key]
            if not art["title"] or not ch.title or ch.section:
                art["title"] = ex["title"] or ch.title or ch.url
            if art["status"] in ("ok", "warn"):
                if art["content_hash"] in seen_hash:
                    other = seen_hash[art["content_hash"]]
                    art.update(status="duplicate", blocks=[], error={"kind": "duplicate",
                               "message": f"Content identical to chapter {other}.", "hint": ""})
                else:
                    seen_hash[art["content_hash"]] = ch.index
                    attach_images(art, fetcher, assets, s, local_dir, budget)
                    art["counts"] = count_blocks(art["blocks"])
        except FetchError as e:
            art["status"] = "failed"
            art["error"] = {"kind": e.kind, "message": e.message, "hint": e.hint}
        except (OSError, UrlError) as e:
            art["status"] = "failed"
            art["error"] = {"kind": "io", "message": str(e), "hint": ""}
        art["title"] = art["title"] or ch.title or f"Chapter {ch.index}"
        flag = {"ok": "ok", "warn": "WARN", "failed": "FAILED", "duplicate": "duplicate"}[art["status"]]
        extra = f" ({art['error']['message']})" if art["error"] else f" conf={art['confidence']}"
        say(f"      -> {flag}{extra}")
        articles.append(art)
    return articles


def report_rows(articles: list[dict]) -> list[dict]:
    rows = []
    for a in articles:
        notes = list(a["warnings"])
        if a["error"]:
            notes.insert(0, a["error"]["message"])
        st = {"ok": "ok", "warn": "warn", "failed": "failed", "duplicate": "skipped"}[a["status"]]
        rows.append({"index": a["index"], "title": a["title"], "section": a.get("section", ""), "url": a["url"], "status": st,
                     "confidence": a["confidence"], "counts": a["counts"], "notes": notes,
                     "content_hash": a["content_hash"], "words": a["words"],
                     "hint": (a["error"] or {}).get("hint", "")})
    return rows


def determine_status(articles: list[dict], checks: list[dict] | None) -> str:
    failed = [a for a in articles if a["status"] == "failed"]
    if failed:
        return "Partial"
    informational = ("Short article",)
    warn = any(a["status"] == "warn" or [w for w in a["warnings"] if not w.startswith(informational)] for a in articles) or \
        any(a["status"] == "duplicate" for a in articles) or (checks and any(not c["ok"] for c in checks))
    return "Completed with warnings" if warn else "Completed"


def build_pipeline(chapters: list[Chapter], s: Settings, fetcher: Fetcher | None, local_mode: bool = False,
                   base_url: str | None = None, source_index_url: str | None = None) -> int:
    if not chapters:
        say("No chapters to build.")
        return 1
    title = s.title or (chapters[0].title if len(chapters) == 1 else "Tutorial Compilation")
    slug_dir = slugify(title)
    out = safe_join(s.out_dir, slug_dir)
    out.mkdir(parents=True, exist_ok=True)
    assets = out / "assets"
    fh = logging.FileHandler(out / "build.log", mode="w", encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logging.getLogger().addHandler(fh)
    try:
        if not local_mode:
            hosts = sorted({urllib.parse.urlsplit(c.url).hostname or "" for c in chapters})
            for h in hosts:
                if not host_acknowledged(s, h):
                    return 1
        articles = process_chapters(chapters, fetcher, s, assets, local_mode, base_url)
        ok_n = sum(a["status"] in ("ok", "warn") for a in articles)
        failed = [a for a in articles if a["status"] == "failed"]
        if ok_n == 0:
            say("\nNo chapter could be extracted; nothing to typeset. See completeness.json.")
            write_reports(out, articles, s, title, "Failed", None, 0, source_index_url)
            return 1
        if failed:
            say(f"\n{len(failed)} of {len(articles)} chapter(s) failed:")
            for a in failed:
                say(f"  #{a['index']} {a['title'][:60]}: {a['error']['message']}")
            proceed = s.allow_partial
            if not proceed and sys.stdin.isatty():
                proceed = input("Generate a PARTIAL document without them? [y/N] ").strip().lower() == "y"
            if not proceed:
                write_reports(out, articles, s, title, "Partial (not generated)", None, 0, source_index_url)
                say("Not generating a partial document (use --allow-partial to confirm). Report written.")
                return 3
        # status is provisional until the outputs have been verified
        status = determine_status(articles, None)
        meta = {"title": title, "date": datetime.now().strftime("%Y-%m-%d")}
        stem = slugify(title, keep_case=True)
        suffix = "-PARTIAL" if failed else ""
        pdf_name, docx_name = f"{stem}{suffix}.pdf", f"{stem}{suffix}.docx"
        pdf_path, docx_path = safe_join(out, pdf_name), safe_join(out, docx_name)
        rows = report_rows(articles)
        doc_articles = [a for a in articles if a["status"] != "duplicate"]     # duplicates: report only

        def make_outputs(st: str):
            checks_, wrapped_, html_ = [], 0, ""
            if s.pdf:
                html_, rend = build_html(doc_articles, s, meta, st, rows)
                say("\nTypesetting PDF ...")
                render_pdf(html_, pdf_path, out)
                checks_ += verify_pdf(pdf_path, doc_articles, meta, s)
                wrapped_ = rend.wrapped_lines
            if s.docx:
                say("\nBuilding Word document ...")
                build_docx(doc_articles, s, meta, st, rows, assets, docx_path)
                checks_ += verify_docx(docx_path, doc_articles, meta, s)
            return checks_, wrapped_, html_

        checks, wrapped, html_doc = make_outputs(status)
        final = determine_status(articles, checks)
        if final != status:                                     # re-make so the files' own report shows the final status
            checks, wrapped, html_doc = make_outputs(final)
            final = determine_status(articles, checks)
        if s.write_markdown:
            (out / f"{stem}.md").write_text(build_markdown(doc_articles, meta, s, final), encoding="utf-8")
        if s.write_html and html_doc:
            (out / f"{stem}.html").write_text(html_doc, encoding="utf-8")
        files = {"pdf": pdf_name if s.pdf else None, "docx": docx_name if s.docx else None}
        write_reports(out, articles, s, title, final, checks, wrapped, source_index_url, files)
        say(f"\n{final.upper()}: {ok_n}/{len(articles)} chapters included, "
            f"{sum(not c['ok'] for c in checks)} verification check(s) failed.")
        for c in checks:
            say(f"  [{'pass' if c['ok'] else 'FAIL'}] {c['name']}" + (f" - {c['detail']}" if c["detail"] else ""))
        say(f"Output: {out}")
        return 0
    finally:
        logging.getLogger().removeHandler(fh)
        fh.close()


def write_reports(out: Path, articles: list[dict], s: Settings, title: str, status: str, checks, wrapped: int,
                  source_index_url: str | None, files: dict | None = None) -> None:
    rows = report_rows(articles)
    summary = {"requested": len(articles), "included": sum(a["status"] in ("ok", "warn") for a in articles),
               "failed": sum(a["status"] == "failed" for a in articles),
               "skipped_duplicates": sum(a["status"] == "duplicate" for a in articles)}
    comp = {"title": title, "status": status, "generated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "summary": summary, "verification": checks, "chapters": rows}
    (out / "completeness.json").write_text(json.dumps(comp, indent=2, ensure_ascii=False), encoding="utf-8")
    md = [f"# Completeness report: {title}", "", f"**Status:** {status}", "",
          f"Requested {summary['requested']}, included {summary['included']}, failed {summary['failed']}, "
          f"skipped duplicates {summary['skipped_duplicates']}.", "",
          "| # | Status | Conf. | Code | Tables | Chapter | Notes |", "|---|---|---|---|---|---|---|"]
    for r in rows:
        md.append(f"| {r['index']} | {r['status']} | {r['confidence']} | {r['counts'].get('code_blocks', 0)} | "
                  f"{r['counts'].get('tables', 0)} | {r['title'][:60].replace('|', '/')} | "
                  f"{'; '.join(r['notes'])[:200].replace('|', '/')} |")
    if checks:
        md += ["", "## PDF verification", ""] + [f"- {'PASS' if c['ok'] else 'FAIL'}: {c['name']} {c['detail']}" for c in checks]
    (out / "completeness.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    manifest = {
        "tool": "TutorialStudyForge", "version": __version__, "schema_version": SCHEMA_VERSION,
        "generated": comp["generated"], "title": title, "status": status, "files": files or {},
        "source_index_url": source_index_url, "settings": s.public(), "code_lines_wrapped": wrapped,
        "chapters": [{k: a[k] for k in ("index", "url", "canonical_url", "final_url", "title", "section", "source_name", "depth",
                                         "status", "confidence", "method", "fetched_at", "from_cache", "content_hash",
                                         "words", "counts", "warnings")} | {"error": a["error"]} for a in articles],
    }
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "chapters.json").write_text(json.dumps(
        {"schema_version": SCHEMA_VERSION, "articles": [{k: v for k, v in a.items()} for a in articles]},
        ensure_ascii=False), encoding="utf-8")


# ======================================================================================
# Local import
# ======================================================================================
def natural_key(p: Path):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", p.name)]


def chapters_from_html_paths(paths: list[str]) -> list[Chapter]:
    files: list[Path] = []
    for raw in paths:
        p = Path(raw)
        if p.is_dir():
            files += sorted([x for x in p.iterdir() if x.suffix.lower() in (".html", ".htm")], key=natural_key)
        elif p.is_file():
            files.append(p)
        else:
            say(f"Not found: {raw}")
            sys.exit(2)
    chapters = []
    for i, f in enumerate(files, 1):
        chapters.append(Chapter(i, f"file:{f.name}", "", 0, source="import", path=f))
    return chapters


# ======================================================================================
# Self-test sample
# ======================================================================================
SELFTEST_PAGES = {
    "01-intro.html": """<html><head><title>Intro - SampleSite</title></head><body><nav>Home | Courses</nav>
<div class="ad-banner">Buy our course now! <a href="x">Sponsored link</a></div>
<article><h1>Introduction to SQL</h1>
<p>SQL (Structured Query Language) is used to <strong>query</strong> and manage relational databases. This chapter introduces
the core ideas and shows a first query so that you can see how statements are structured and how results are returned.</p>
<h2>Your first query</h2><p>Use <code>SELECT</code> to read rows from a table. The following example reads two columns:</p>
<pre class="language-sql"><code>SELECT name,
       salary
  FROM employees
 WHERE salary &gt; 50000
 ORDER BY salary DESC;</code></pre>
<div class="note"><p>Note: keywords are case-insensitive, but upper case is a common convention.</p></div>
<h2>Why it matters</h2><ul><li>Declarative: you say <em>what</em>, not how<ul><li>The engine plans the work</li></ul></li><li>Portable across systems</li></ul>
<div class="share-buttons"><a href="#">Share</a><a href="#">Tweet</a></div></article></body></html>""",
    "02-joins.html": """<html><head><title>Joins</title></head><body><main><h1>SQL Joins</h1>
<p>A join combines rows from two tables. The table below summarises the common join types and what each one returns when
rows have no match, which is the most frequent source of confusion for beginners learning relational queries.</p>
<table><thead><tr><th>Join</th><th colspan="2">Behaviour</th></tr></thead><tbody>
<tr><td>INNER</td><td>Matches only</td><td>Common rows</td></tr><tr><td>LEFT</td><td>All left rows</td><td>NULL when no match</td></tr></tbody></table>
<div class="highlight"><table><tr><td class="linenos"><pre>1
2
3</pre></td><td class="code"><pre class="language-sql">SELECT e.name, d.title
FROM employees e
    JOIN departments d ON d.id = e.dept_id;</pre></td></tr></table></div>
<p>Output of the query above is a list of employee names with their department titles, one row per employee.</p>
<pre>def very_long(a, b):
    return some_function_with_a_long_name(argument_number_one, argument_number_two, argument_number_three, argument_number_four)</pre>
</main></body></html>""",
    "03-python.html": """<html><head><title>Python Basics</title></head><body><article><h1>Python Basics</h1>
<p>Python uses indentation to define blocks. Indentation is part of the syntax, which is why preserving whitespace in code
examples is essential when compiling study material from the web into a document that people will read offline.</p>
<pre><code class="language-python">def greet(name):
    if name:
        print(f"Hello, {name}!")
    else:
        print("Hello, world!")</code></pre>
<p>Run it by calling <code>greet("Ada")</code>.</p></article></body></html>""",
}


def cmd_selftest(s: Settings) -> int:
    d = s.out_dir / "_selftest_input"
    d.mkdir(parents=True, exist_ok=True)
    for name, content in SELFTEST_PAGES.items():
        (d / name).write_text(content, encoding="utf-8")
    s.title = s.title or "Selftest Sample Guide"
    s.write_markdown = True
    s.docx = s.docx or HAVE_DOCX
    chapters = chapters_from_html_paths([str(d)])
    return build_pipeline(chapters, s, None, local_mode=True)


# ======================================================================================
# inspect
# ======================================================================================
def cmd_sites() -> int:
    say("Built-in rules (written from each site's usual layout; sites change, so check with 'discover' / 'inspect'):\n")
    for key, prof in DEFAULT_PROFILES.items():
        if key == "*":
            continue
        say(f"  {prof.get('name', key):<26} {prof.get('example', '')}")
    say("\nOther sites usually work through the generic detectors: article/main detection, side-menu detection and "
        "heading-based sections.\nFix anything that is off with --selector, --link-selector, --section-selector "
        "or --profiles (see README).")
    return 0


def cmd_inspect(target: str, s: Settings, fetcher: Fetcher | None) -> int:
    if re.match(r"^https?://", target):
        assert fetcher is not None
        host = urllib.parse.urlsplit(target).hostname or ""
        if not host_acknowledged(s, host):
            return 1
        res = fetcher.get(target)
        data, url, base = res.body, res.final_url, res.final_url
    else:
        data, url, base = Path(target).read_bytes(), target, None
    art = extract_article(data, url, base, s, profile_host=urllib.parse.urlsplit(url).hostname or "")
    say(f"Title:       {art['title']!r}")
    say(f"Container:   {art['method']}")
    say(f"Status:      {art['status']}  confidence={art['confidence']}  words={art['words']}")
    say(f"Counts:      {art['counts']}")
    if art["error"]:
        say(f"Error:       {art['error']['message']}\n             {art['error']['hint']}")
    for w in art["warnings"]:
        say(f"Warning:     {w}")
    say("\nOutline:")
    for b in art["blocks"]:
        if b["type"] == "heading":
            say(f"  {'  ' * (b['level'] - 2)}h{b['level']} {re.sub('<[^>]+>', '', htmllib.unescape(b['html']))[:80]}")
    codes = [b for b in art["blocks"] if b["type"] == "code"]
    for i, b in enumerate(codes[:3], 1):
        say(f"\nCode block {i} (lang={b['language'] or '?'}):")
        for ln in b["text"].split("\n")[:8]:
            say(f"  |{ln}")
    return 0 if art["status"] in ("ok", "warn") else 1


# ======================================================================================
# DOCX output: an editable Word document built from the same block model
# ======================================================================================
try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Mm, Pt, RGBColor
    HAVE_DOCX = True
except ImportError:                                     # reported when a .docx is requested
    HAVE_DOCX = False

_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def xml_safe(s: str) -> str:
    return _XML_BAD.sub("", s)


def html_plain(h: str) -> str:
    return collapse_ws(re.sub(r"<[^>]+>", "", htmllib.unescape(h)))


class DocxBuilder:
    """Builds the Word file with real Word styles (Heading 1-9, List Paragraph + numbering, Table Grid, Code,
    Callout) so that the result is easy to edit, restyle and re-flow."""

    BODY_FONT, CODE_FONT = "Calibri", "Consolas"

    def __init__(self, s: Settings, assets: Path):
        self.s, self.assets = s, assets
        self.doc = Document()
        self.hshift = 0
        self._page()
        self._styles()
        self._numbering()

    # ---- setup ----
    def _page(self) -> None:
        s = self.s
        w_mm, h_mm = {"A4": (210.0, 297.0), "LETTER": (215.9, 279.4)}.get(s.page_size.upper(), (210.0, 297.0))
        sec = self.doc.sections[0]
        if s.landscape:
            w_mm, h_mm = h_mm, w_mm
            sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Mm(w_mm), Mm(h_mm)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(s.margin_mm)
        self.content_w = Mm(w_mm - 2 * s.margin_mm)
        sec.different_first_page_header_footer = True            # clean cover page

    @staticmethod
    def _bdr_shd(style, fill: str, left: str | None = None, box: str | None = None) -> None:
        pPr = style.element.get_or_add_pPr()
        parts = []
        if box:
            parts += [f'<w:top w:val="single" w:sz="4" w:space="3" w:color="{box}"/>']
        parts += [f'<w:left w:val="single" w:sz="{18 if left else 4}" w:space="6" w:color="{left or box or "BBBBBB"}"/>']
        if box:
            parts += [f'<w:bottom w:val="single" w:sz="4" w:space="3" w:color="{box}"/>',
                      f'<w:right w:val="single" w:sz="4" w:space="6" w:color="{box}"/>']
        pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}>{"".join(parts)}</w:pBdr>'))
        pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))

    def _new(self, name: str, kind=WD_STYLE_TYPE.PARAGRAPH, base: str | None = "Normal"):
        try:
            return self.doc.styles[name]
        except KeyError:
            st = self.doc.styles.add_style(name, kind)
            if base and kind == WD_STYLE_TYPE.PARAGRAPH:
                st.base_style = self.doc.styles[base]
            return st

    def _styles(self) -> None:
        s, st = self.s, self.doc.styles
        dark = RGBColor(0x1B, 0x2A, 0x41)
        n = st["Normal"]
        n.font.name, n.font.size = self.BODY_FONT, Pt(s.font_pt)
        n.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.BODY_FONT)
        n.paragraph_format.space_after, n.paragraph_format.line_spacing = Pt(6), 1.12
        for name, size in (("Heading 1", 22), ("Heading 2", 17), ("Heading 3", 14), ("Heading 4", 12), ("Heading 5", 11),
                           ("Heading 6", 11), ("Heading 7", 10.5), ("Heading 8", 10.5), ("Heading 9", 10.5)):
            h = st[name]
            h.font.name, h.font.size, h.font.bold, h.font.color.rgb = self.BODY_FONT, Pt(size), True, dark
            h.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.BODY_FONT)
            h.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
            h.paragraph_format.space_after, h.paragraph_format.keep_with_next = Pt(6), True
        for name, size in (("Title", 30), ("Subtitle", 15)):
            x = st[name]
            x.font.name, x.font.size, x.font.color.rgb = self.BODY_FONT, Pt(size), dark if name == "Title" else RGBColor(0x44, 0x44, 0x44)
            x.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # code block
        code = self._new("Code")
        self._bdr_shd(code, "F5F5F5", left="444444", box="D0D0D0")
        code.font.name, code.font.size = self.CODE_FONT, Pt(max(8.0, s.font_pt - 1.5))
        rf = code.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:cs"), "Courier New"); rf.set(qn("w:eastAsia"), self.CODE_FONT)
        pf = code.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(8), 1.0
        pf.left_indent = pf.right_indent = Mm(1.5)
        lab = self._new("Code Label")
        lab.font.size, lab.font.bold, lab.font.color.rgb = Pt(8), True, RGBColor(0x55, 0x55, 0x55)
        lab.paragraph_format.space_after, lab.paragraph_format.keep_with_next = Pt(1), True
        # inline code
        ic = self._new("Inline Code", WD_STYLE_TYPE.CHARACTER)
        ic.font.name, ic.font.size = self.CODE_FONT, Pt(max(8.0, s.font_pt - 1))
        ic.element.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="EFEFEF"/>'))
        hl = self._new("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        hl.font.color.rgb, hl.font.underline = RGBColor(0x0B, 0x4A, 0x8F), True
        # callout / note
        co = self._new("Callout")
        self._bdr_shd(co, "F7F7F7", left="888888")
        co.paragraph_format.left_indent, co.paragraph_format.space_after = Mm(3), Pt(6)
        src = self._new("Source Line")
        src.font.size, src.font.color.rgb = Pt(8.5), RGBColor(0x66, 0x66, 0x66)
        src.paragraph_format.space_after = Pt(10)
        part = self._new("Part Description")
        part.font.italic, part.font.size = True, Pt(s.font_pt + 1.5)
        part.paragraph_format.space_after = Pt(12)
        ct = self._new("Contents Title")
        ct.font.name, ct.font.size, ct.font.bold, ct.font.color.rgb = self.BODY_FONT, Pt(22), True, dark
        ct.paragraph_format.space_after = Pt(12)
        for lvl, ind in ((1, 0), (2, 6), (3, 12)):
            tc = self._new(f"toc {lvl}")
            tc.paragraph_format.left_indent, tc.paragraph_format.space_after = Mm(ind), Pt(2)
            tc.font.bold = (lvl == 1) and None
        try:
            st["Caption"].font.size = Pt(9)
        except KeyError:
            pass

    def _numbering(self) -> None:
        root = self.doc.part.numbering_part.element
        nums = [int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))]
        abss = [int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))]
        self._next_num = max(nums + [0]) + 1
        base = max(abss + [0]) + 1
        self.abs_bullet, self.abs_decimal = base, base + 1
        first_num = root.find(qn("w:num"))
        for aid, kind in ((self.abs_bullet, "bullet"), (self.abs_decimal, "decimal")):
            el = parse_xml(self._abstract_xml(aid, kind))
            first_num.addprevious(el) if first_num is not None else root.append(el)
        self.bullet_num = self._new_num(self.abs_bullet)

    @staticmethod
    def _abstract_xml(aid: int, kind: str) -> str:
        lv = []
        for i in range(6):
            if kind == "bullet":
                fmt, txt = "bullet", ["•", "◦", "▪"][i % 3]
            else:
                fmt, txt = ["decimal", "lowerLetter", "lowerRoman"][i % 3], f"%{i + 1}."
            lv.append(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{txt}"/>'
                      f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{540 + 360 * i}" w:hanging="360"/></w:pPr></w:lvl>')
        return (f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="hybridMultilevel"/>'
                f'{"".join(lv)}</w:abstractNum>')

    def _new_num(self, abs_id: int, start: int | None = None) -> int:
        n, self._next_num = self._next_num, self._next_num + 1
        ov = (f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{start}"/></w:lvlOverride>' if start else "")
        self.doc.part.numbering_part.element.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{n}"><w:abstractNumId w:val="{abs_id}"/>{ov}</w:num>'))
        return n

    # ---- inline content ----
    def par(self, style: str = "Normal"):
        return self.doc.add_paragraph(style=style)

    def inline(self, p, html_str: str, bold: bool = False) -> None:
        root = BeautifulSoup(f"<x>{html_str}</x>", "html.parser").x
        self._inline(p, root, {"b": bold, "i": False, "code": False, "sub": False, "sup": False, "href": None})

    def _inline(self, p, node, f: dict) -> None:
        for ch in node.children:
            if isinstance(ch, NavigableString):
                if str(ch):
                    self._run(p, str(ch), f)
            elif isinstance(ch, Tag):
                g = dict(f)
                n = ch.name
                if n == "br":
                    p.add_run().add_break()
                    continue
                if n == "strong":
                    g["b"] = True
                elif n == "em":
                    g["i"] = True
                elif n == "code":
                    g["code"] = True
                elif n in ("sub", "sup"):
                    g[n] = True
                elif n == "a" and ch.get("href"):
                    g["href"] = ch["href"]
                self._inline(p, ch, g)

    def _run(self, p, text: str, f: dict) -> None:
        text = xml_safe(text)
        if f["href"]:
            try:
                rid = p.part.relate_to(f["href"], RT.HYPERLINK, is_external=True)
            except Exception:
                rid = None
            if rid:
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), rid)
                r = OxmlElement("w:r")
                rpr = OxmlElement("w:rPr")
                rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "Hyperlink"); rpr.append(rs)
                if f["b"]:
                    rpr.append(OxmlElement("w:b"))
                r.append(rpr)
                tt = OxmlElement("w:t"); tt.text = text; tt.set(qn("xml:space"), "preserve"); r.append(tt)
                h.append(r)
                p._p.append(h)
                return
        r = p.add_run(text)
        if f["b"]:
            r.bold = True
        if f["i"]:
            r.italic = True
        if f["code"]:
            r.style = "Inline Code"
        if f["sub"]:
            r.font.subscript = True
        if f["sup"]:
            r.font.superscript = True

    # ---- blocks ----
    def blocks(self, bs: list[dict], indent_mm: float = 0.0, style: str = "Normal") -> None:
        for b in bs:
            self.block(b, indent_mm, style)

    def block(self, b: dict, indent_mm: float = 0.0, style: str = "Normal") -> None:
        t = b["type"]
        if t == "paragraph":
            p = self.par(style)
            self.inline(p, b["html"])
            self._indent(p, indent_mm)
        elif t == "heading":
            lvl = max(1, min(9, b["level"] + self.hshift))
            self.par(f"Heading {lvl}").add_run(xml_safe(html_plain(b["html"])))
        elif t in ("code", "equation"):
            text = b["text"] if t == "code" else b["tex"]
            lang = b.get("language", "") if t == "code" else "equation"
            if lang:
                self._indent(self.par("Code Label"), indent_mm).add_run(xml_safe(lang.upper()))
            p = self.par("Code")
            self._set_code_text(p, xml_safe(text))
            self._indent(p, indent_mm + 1.5)
        elif t == "list":
            self._list(b, 0)
        elif t == "table":
            self._table(b)
        elif t == "image":
            self._image(b)
        elif t == "quote":
            if b["kind"]:
                self.par("Callout").add_run(b["kind"].title()).bold = True
            for sub in b["blocks"]:
                self.block(sub, indent_mm, "Callout" if sub["type"] == "paragraph" else style)
        elif t == "hr":
            p = self.par()
            p._p.get_or_add_pPr().append(parse_xml(
                f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="999999"/></w:pBdr>'))

    @staticmethod
    def _set_code_text(p, text: str) -> None:
        p.add_run(text)                                   # newline -> line break, tab -> tab (python-docx)
        for tnode in p._p.iter(qn("w:t")):
            tnode.set(qn("xml:space"), "preserve")        # keep leading / repeated spaces exactly

    @staticmethod
    def _indent(p, mm: float):
        if mm:
            p.paragraph_format.left_indent = Mm(mm)
        return p

    def _list(self, b: dict, level: int) -> None:
        lvl = min(level, 5)
        num = self._new_num(self.abs_decimal, b["start"]) if b["ordered"] else self.bullet_num
        cont_mm = (540 + 360 * lvl) * 25.4 / 1440
        for item in b["items"]:
            first = True
            for blk in item:
                if first and blk["type"] == "paragraph":
                    p = self.par("List Paragraph")
                    ppr = p._p.get_or_add_pPr()
                    npr = ppr.get_or_add_numPr()
                    npr.get_or_add_ilvl().val = lvl
                    npr.get_or_add_numId().val = num
                    self.inline(p, blk["html"])
                elif blk["type"] == "list":
                    self._list(blk, level + 1)
                else:
                    self.block(blk, cont_mm)
                first = False

    def _table(self, b: dict) -> None:
        rows = b["rows"]
        occupied: set[tuple[int, int]] = set()
        placed = []
        for ri, row in enumerate(rows):
            ci = 0
            for cell in row["cells"]:
                while (ri, ci) in occupied:
                    ci += 1
                rs = min(cell["rowspan"], len(rows) - ri)
                placed.append((ri, ci, rs, cell["colspan"], cell, row["head"]))
                for dr in range(rs):
                    for dc in range(cell["colspan"]):
                        occupied.add((ri + dr, ci + dc))
                ci += cell["colspan"]
        ncols = max(c + cs for _, c, _, cs, _, _ in placed)
        if b["caption"]:
            self.par("Caption").add_run(xml_safe(b["caption"]))
        tbl = self.doc.add_table(rows=len(rows), cols=ncols)
        tbl.style = "Table Grid"
        tw = tbl._tbl.tblPr.find(qn("w:tblW"))
        if tw is not None:                                    # full text width
            tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
        for ri, ci, rs, cs, cell, head in placed:
            c = tbl.cell(ri, ci)
            if rs > 1 or cs > 1:
                c = c.merge(tbl.cell(ri + rs - 1, ci + cs - 1))
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            self.inline(p, cell["html"], bold=head or cell["th"])
            for r in p.runs:
                r.font.size = Pt(max(8.0, self.s.font_pt - 1))
            if head or cell["th"]:
                c._tc.get_or_add_tcPr().append(parse_xml(
                    f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="E8E8E8"/>'))
        for ri, row in enumerate(rows):
            trpr = tbl.rows[ri]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if row["head"]:
                trpr.append(OxmlElement("w:tblHeader"))
        self.par().paragraph_format.space_after = Pt(2)

    def _webp_to_png(self, path: Path) -> Path:
        from PIL import Image
        out = path.with_suffix(".png")
        if not out.exists():
            Image.open(path).convert("RGBA").save(out)
        return out

    def _image(self, b: dict) -> None:
        if not b.get("asset"):
            if not b.get("skipped"):
                p = self.par()
                p.add_run(f"[Image not available: {xml_safe(b['caption'] or b['alt'] or 'image')}]").italic = True
            return
        path = self.assets / b["asset"]
        try:
            if path.suffix.lower() == ".webp":
                path = self._webp_to_png(path)
            from docx.image.image import Image as DImage
            im = DImage.from_file(str(path))
            dpi = im.horz_dpi or 96
            w = min(Mm(im.px_width / dpi * 25.4), int(self.content_w * 0.92), Mm(170))
            if im.px_width and im.px_height / im.px_width * w > Mm(190):
                w = int(Mm(190) * im.px_width / im.px_height)
            p = self.par()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=w)
        except Exception as e:
            p = self.par()
            p.add_run(f"[Image could not be embedded ({e.__class__.__name__}): {b.get('url') or b['src']}]").italic = True
            return
        if b["caption"]:
            cp = self.par("Caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.add_run(xml_safe(b["caption"]))

    # ---- document parts ----
    def _field(self, p, instr: str, placeholder: str = "1") -> None:
        for kind, text in (("begin", None), (None, instr), ("separate", None), (None, placeholder), ("end", None)):
            r = p.add_run()
            if kind:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
            elif text == instr:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
            else:
                r.text = text

    def header_footer(self, title: str) -> None:
        sec = self.doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(xml_safe(title[:80]))
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x66, 0x66, 0x66)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE")

    def cover(self, meta: dict, status: str, n_ok: int, n_parts: int) -> None:
        s = self.s
        sp = self.par(); sp.paragraph_format.space_before = Pt(150)
        self.par("Title").add_run(xml_safe(meta["title"]))
        if s.subtitle:
            self.par("Subtitle").add_run(xml_safe(s.subtitle))
        lines = ([s.author] if s.author else []) + [f"Compiled {meta['date']}",
                 f"{n_ok} chapters" + (f" in {n_parts} parts" if n_parts else "")]
        for ln in lines:
            p = self.par(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(xml_safe(ln)); r.font.size, r.font.color.rgb = Pt(10.5), RGBColor(0x44, 0x44, 0x44)
        p = self.par(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(30)
        r = p.add_run("Compiled for personal study from the sources listed at the end of this document. Content belongs "
                      "to its original authors; do not redistribute. Faithful compilation: text was not rewritten or "
                      "summarised.")
        r.font.size, r.italic = Pt(8.5), True
        if status == "Partial":
            p = self.par(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("PARTIAL DOCUMENT: some chapters could not be retrieved (see the build report)")
            r.bold, r.font.color.rgb = True, RGBColor(0x99, 0, 0)

    def contents(self, entries: list[tuple[int, str]], depth: int) -> None:
        h = self.par("Contents Title")
        h.paragraph_format.page_break_before = True
        h.add_run("Contents")
        instr = f'TOC \\o "1-{depth}" \\h \\z \\u'
        entries = entries or [(1, "Right-click here and choose Update Field to build the contents")]
        for i, (lvl, text) in enumerate(entries):
            p = self.par(f"toc {min(lvl, 3)}")
            if i == 0:
                for kind in ("begin", None, "separate"):
                    r = p.add_run()
                    if kind:
                        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
                    else:
                        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
            p.add_run(xml_safe(text))
            if i == len(entries) - 1:
                r = p.add_run()
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
        # ask Word to refresh the contents (with page numbers) when the file is opened
        st = self.doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)

    def heading(self, level: int, text: str, page_break: bool):
        p = self.par(f"Heading {level}")
        p.add_run(xml_safe(text))
        if page_break:
            p.paragraph_format.page_break_before = True
        return p

    def source_line(self, url: str) -> None:
        p = self.par("Source Line")
        p.add_run("Source: ")
        self._run(p, url, {"b": False, "i": False, "code": False, "sub": False, "sup": False, "href": url})


def build_docx(articles: list[dict], s: Settings, meta: dict, status: str, report_rows: list[dict],
               assets: Path, out_path: Path) -> None:
    if not HAVE_DOCX:
        say("Missing dependency 'python-docx' (needed for Word output). Install with: pip install python-docx")
        sys.exit(1)
    d = DocxBuilder(s, assets)
    has_parts = any(a.get("section") for a in articles)
    included = [a for a in articles if a["status"] in ("ok", "warn")]
    n_parts = len({a["section"] for a in articles if a.get("section")})
    cp = d.doc.core_properties
    cp.title, cp.subject, cp.author = meta["title"], s.subtitle or "", s.author or APP_NAME
    cp.last_modified_by, cp.comments = APP_NAME, f"Compiled by {APP_NAME} (TSF) {__version__}; status: {status}"
    d.header_footer(meta["title"])
    if s.cover:
        d.cover(meta, status, len(included), n_parts if has_parts else 0)
    if s.toc:
        entries, seen = [], None
        for a in articles:
            sec = a.get("section") or ""
            if has_parts and sec and sec != seen:
                seen = sec
                entries.append((1, sec))
            entries.append((2 if (has_parts and sec) else 1, a["title"] or f"Chapter {a['index']}"))
        toc_depth = s.toc_depth if s.toc_depth is not None else (1 if has_parts else 2)
        d.contents(entries, (toc_depth + 1) if has_parts else toc_depth)
    cur = None
    for a in articles:
        sec = a.get("section") or ""
        in_part = has_parts and bool(sec)
        if in_part and sec != cur:
            cur = sec
            d.heading(1, sec, True)
            if a.get("section_desc"):
                d.par("Part Description").add_run(xml_safe(a["section_desc"]))
        d.heading(2 if in_part else 1, a["title"] or f"Chapter {a['index']}", s.chapter_breaks)
        d.hshift = 1 if in_part else 0
        if a["status"] in ("ok", "warn"):
            if s.source_urls and a["url"].startswith("http"):
                d.source_line(a["url"])
            d.blocks(a["blocks"])
        else:
            err = a.get("error") or {}
            p = d.par("Callout")
            r = p.add_run("This chapter is missing from this document.")
            r.bold, r.font.color.rgb = True, RGBColor(0x99, 0, 0)
            p.add_run(f" Reason: {xml_safe(err.get('message', a['status']))}. Source: {xml_safe(a['url'])}")
    d.hshift = 0
    if s.source_urls:
        links = [a for a in articles if a["url"].startswith("http")]
        if links:
            d.heading(1, "Sources", True)
            d.par().add_run("Original locations of the compiled chapters.")
            num = d._new_num(d.abs_bullet)
            for a in links:
                p = d.par("List Paragraph")
                npr = p._p.get_or_add_pPr().get_or_add_numPr()
                npr.get_or_add_ilvl().val = 0
                npr.get_or_add_numId().val = num
                p.add_run(xml_safe((a["title"] or a["url"]) + ": "))
                d._run(p, a["url"], {"b": False, "i": False, "code": False, "sub": False, "sup": False, "href": a["url"]})
    if s.report_page:
        d.heading(1, "Build report", True)
        d.par().add_run(f"Status: {status}. Generated by {APP_NAME} (TSF) {__version__}.")
        tb = {"type": "table", "caption": "", "rows": [
            {"head": True, "cells": [{"html": h, "th": True, "colspan": 1, "rowspan": 1}
                                     for h in ("#", "Chapter", "Status", "Conf.", "Code", "Tables", "Notes")]}] + [
            {"head": False, "cells": [{"html": esc(str(v)), "th": False, "colspan": 1, "rowspan": 1} for v in (
                x["index"], x["title"][:70], x["status"], x["confidence"], x["counts"].get("code_blocks", 0),
                x["counts"].get("tables", 0), "; ".join(x["notes"])[:160])]} for x in report_rows]}
        d._table(tb)
    d.doc.save(str(out_path))


def _walk_blocks(bs: list[dict]):
    for b in bs:
        yield b
        if b["type"] == "list":
            for it in b["items"]:
                yield from _walk_blocks(it)
        elif b["type"] == "quote":
            yield from _walk_blocks(b["blocks"])


def verify_docx(path: Path, articles: list[dict], meta: dict, s: Settings) -> list[dict]:
    """Re-open the finished .docx and check it independently of the code that wrote it."""
    from collections import Counter
    checks: list[dict] = []

    def add(name, ok, detail=""):
        checks.append({"name": name, "ok": bool(ok), "detail": detail})
    try:
        doc = Document(str(path))
    except Exception as e:
        add("Word: file can be opened", False, str(e))
        return checks
    add("Word: file can be opened", path.stat().st_size > 0, f"{path.stat().st_size} bytes")
    paras = doc.paragraphs
    heads = {p.text for p in paras if p.style.name.startswith("Heading")}
    included = [a for a in articles if a["status"] in ("ok", "warn")]
    missing = [a["title"] for a in included if xml_safe(a["title"] or "") not in heads]
    add("Word: a heading for every chapter", not missing, f"missing: {missing[:3]}" if missing else f"{len(included)} chapters")
    secs = sorted({a["section"] for a in included if a.get("section")})
    if secs:
        miss = [x for x in secs if xml_safe(x) not in heads]
        add("Word: a heading for every part", not miss, f"missing: {miss[:3]}" if miss else f"{len(secs)} parts")
    expected = Counter()
    n_images = 0
    for a in included:
        for b in _walk_blocks(a["blocks"]):
            if b["type"] == "code":
                expected[xml_safe(b["text"])] += 1
            elif b["type"] == "equation":
                expected[xml_safe(b["tex"])] += 1
            elif b["type"] == "image" and b.get("asset"):
                n_images += 1
    found = Counter(p.text for p in paras if p.style.name == "Code")
    lost = expected - found
    add("Word: code blocks identical to the source", not lost,
        f"{sum(expected.values()) - sum(lost.values())}/{sum(expected.values())} exact")
    body_xml = doc.element.xml
    if s.toc:
        add("Word: contents field present", "TOC \\o" in body_xml)
    add("Word: images embedded", len(doc.inline_shapes) == n_images, f"{len(doc.inline_shapes)}/{n_images}")
    return checks


# ======================================================================================
# CLI
# ======================================================================================
COMMANDS = {"go", "discover", "build", "inspect", "selftest", "sites"}

EXAMPLES = """
examples:
  python3 tsf.py                                  guided mode: asks you step by step
  python3 tsf.py https://www.geeksforgeeks.org/sql-tutorial/
                                                           guided mode for that tutorial
  python3 tsf.py go URL --yes --docx --markdown   no questions: everything, as PDF + Word + Markdown
  python3 tsf.py build --index URL --docx --no-pdf  Word document only
  python3 tsf.py sites                                     sites with built-in rules
  python3 tsf.py build --index URL --sections 1-5 only sections 1 to 5
  python3 tsf.py discover URL -o topics.txt       write an editable topic list, then:
  python3 tsf.py build --chapters topics.txt
  python3 tsf.py inspect URL                      show how one page is read (for debugging)
"""


# ---- friendly argument handling ---------------------------------------------------
def fix_url(v: str) -> str:
    """Accept pasted addresses in messy forms; explain clearly when it is not an address at all."""
    v = (v or "").strip().strip("<>\"'")
    m = re.match(r"^\[[^\]]*\]\((\S+?)\)$", v)          # [text](https://...) pasted from chat apps
    if m:
        v = m.group(1)
    if re.match(r"^https?://", v, re.I):
        return v
    if re.fullmatch(r"(?:www\.)?[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+(?:[/?#]\S*)?", v):
        return "https://" + v
    raise UrlError(f"'{v}' is not a web address. Open the tutorial in your browser, copy the address from the "
                   f"address bar and paste it, e.g. https://www.geeksforgeeks.org/sql-tutorial/")


def fix_host(v: str) -> str:
    v = (v or "").strip().strip("<>\"'")
    m = re.match(r"^\[([^\]]*)\]\(.*\)$", v)
    if m:
        v = m.group(1)
    if "://" in v:
        v = urllib.parse.urlsplit(v).hostname or v
    return v.lower()


def ask(prompt: str, default: str = "") -> str:
    try:
        v = input(prompt).strip()
    except EOFError:
        return default
    return v or default


def yesno(prompt: str, default: bool = False) -> bool:
    v = ask(f"{prompt} [{'Y/n' if default else 'y/N'}] ").lower()
    return default if not v else v.startswith("y")


def apply_filters(chapters: list[Chapter], a: argparse.Namespace, s: Settings) -> list[Chapter]:
    if getattr(a, "select", None):
        keep = parse_selection(a.select, len(chapters))
        chapters = [c for c in chapters if c.index in keep]
    if getattr(a, "include", None):
        rx = re.compile(a.include, re.I)
        chapters = [c for c in chapters if rx.search(c.url) or rx.search(c.title)]
    if getattr(a, "exclude", None):
        rx = re.compile(a.exclude, re.I)
        chapters = [c for c in chapters if not (rx.search(c.url) or rx.search(c.title))]
    if len(chapters) > s.max_pages:
        say(f"Limiting to the first {s.max_pages} chapters (--max-pages).")
        chapters = chapters[:s.max_pages]
    for n, c in enumerate(chapters, 1):
        c.index = n
    return chapters


def print_sections(chapters: list[Chapter], info: dict, host: str) -> None:
    secs = sections_of(chapters)
    name = info.get("page_title") or host
    if secs:
        say(f'\nFound {len(chapters)} topics in {len(secs)} sections of "{name}":\n')
        for i, sname in enumerate(secs, 1):
            n = sum(1 for c in chapters if c.section == sname)
            say(f"  {i:>2}. {sname}  ({n} topic{'s' if n != 1 else ''})")
    else:
        say(f'\nFound {len(chapters)} topics in "{name}" (no sections detected).')
        for c in chapters[:12]:
            say(f"  {c.index:>3}. {c.title}")
        if len(chapters) > 12:
            say(f"  ... and {len(chapters) - 12} more")
    for n in info.get("notes", []):
        say(f"  note: {n}")


# ---- guided mode ------------------------------------------------------------------
def run_guided(a: argparse.Namespace, s: Settings, fetcher: Fetcher) -> int:
    interactive = sys.stdin.isatty() and not s.yes
    url = getattr(a, "go_url", None)
    for attempt in range(3):
        if not url:
            if not interactive:
                say("Give the tutorial address, e.g.:\n  python3 tsf.py https://www.geeksforgeeks.org/sql-tutorial/")
                return 2
            if attempt == 0:
                say("\nTutorialStudyForge (TSF): forge an online tutorial into one PDF / Word study guide.\n"
                    "Paste the address of the tutorial's main page (the page that lists all its topics).")
            url = ask("Address: ")
        try:
            url = fix_url(url)
            break
        except UrlError as e:
            say(f"  {e}")
            if not interactive:
                return 2
            url = None
    else:
        return 2
    host = urllib.parse.urlsplit(url).hostname or ""
    if not host_acknowledged(s, host):
        return 1
    say("\nReading the tutorial's topic list ...")
    chapters, info = discover(fetcher, url, s, a.link_selector, a.prefix)
    if not chapters:
        say("No topic links found on that page. Try a different page of the tutorial, or --link-selector "
            "(see 'discover --help').")
        return 1
    print_sections(chapters, info, host)
    secs = sections_of(chapters)
    spec = a.sections
    if secs and not spec and interactive:
        say('\nWhich sections do you want?  Press Enter for ALL, or type numbers such as 1-5,8  '
            '(or "not 12,13" to leave some out).')
        while True:
            spec = ask("Sections: ")
            try:
                filter_sections(chapters, spec)
                break
            except ValueError as e:
                say(f"  {e}")
    chapters = filter_sections(chapters, spec, a.skip_sections)
    chapters = apply_filters(chapters, a, s)
    if not chapters:
        say("Nothing selected.")
        return 1
    default_title = s.title or info.get("page_title") or host
    if interactive and not s.title:
        s.title = ask(f"\nTitle of the PDF [{default_title}]: ", default_title)
    s.title = s.title or default_title
    if not s.subtitle:
        sub = info.get("page_subtitle", "")
        if interactive and sub:
            say(f"Suggested subtitle: {sub}")
            v = ask("Subtitle [Enter = use it, '-' = none, or type your own]: ", sub)
            sub = "" if v == "-" else v
        s.subtitle = sub
    if interactive and not s.docx:
        s.docx = yesno("Also save an editable Word document (.docx)?", True)
    if interactive and not s.write_markdown:
        s.write_markdown = yesno("Also save an editable Markdown copy?")
    folder = safe_join(s.out_dir, slugify(s.title))
    folder.mkdir(parents=True, exist_ok=True)
    topics = folder / "topics.txt"
    write_chapters_file(topics, chapters, info)
    if interactive:
        say(f"\nThe list of topics was saved to: {topics}")
        if ask("Press Enter to start, or type e to edit the topic list first: ").lower().startswith("e"):
            say("Open that file in an editor, delete the lines of topics you do not want, save it, then come back.")
            ask("Press Enter when you are done ... ")
            chapters = read_chapters_file(topics)
            for n, c in enumerate(chapters, 1):
                c.index = n
    est = max(1, round(len(chapters) * max(s.delay, 1.0) * 1.3 / 60))
    kinds = " + ".join(k for k, on in (("PDF", s.pdf), ("Word", s.docx)) if on)
    say(f'\nReady: {len(chapters)} topics -> {kinds} titled "{s.title}".')
    say(f"Estimated time: about {est} minute{'s' if est != 1 else ''} (polite 1-second pauses; pages already downloaded are instant).")
    if interactive and not yesno("Start now?", True):
        say(f"Cancelled. You can start later with:\n  python3 tsf.py build --chapters {topics}")
        return 0
    say("")
    rc = build_pipeline(chapters, s, fetcher, source_index_url=info["index_url"])
    if rc == 0:
        say(f"\nTo rebuild later (for example after editing the list):\n  python3 tsf.py build --chapters {topics}")
    return rc


# ---- parser -----------------------------------------------------------------------
def make_parser() -> argparse.ArgumentParser:
    common = argparse.ArgumentParser(add_help=False)
    g = common.add_argument_group("fetching")
    g.add_argument("--delay", type=float, default=1.0, help="seconds between requests to one site (default 1.0)")
    g.add_argument("--timeout", type=float, default=30.0)
    g.add_argument("--cache-dir", type=Path)
    g.add_argument("--refresh", action="store_true", help="ignore cached pages")
    g.add_argument("--offline", action="store_true", help="use only cached pages")
    g.add_argument("--ack-terms", action="append", default=[], metavar="SITE",
                   help="confirm you reviewed SITE's terms of use (e.g. www.geeksforgeeks.org); asked once if omitted")
    g.add_argument("--allow-private", action="store_true", help="TESTING ONLY: allow localhost/private addresses")
    g2 = common.add_argument_group("extraction")
    g2.add_argument("--profiles", type=Path, help="JSON file of per-site selectors (see README)")
    g2.add_argument("--selector", help="CSS selector of the article body (overrides the built-in rules)")
    g2.add_argument("--drop-text", action="append", default=[], metavar="REGEX",
                    help="drop paragraphs whose whole text matches REGEX (repeatable)")
    g2.add_argument("--min-words", type=int, default=30)
    common.add_argument("-v", "--verbose", action="store_true")

    disc = argparse.ArgumentParser(add_help=False)
    d = disc.add_argument_group("finding topics")
    d.add_argument("--link-selector", help="CSS selector for the topic links (or their container)")
    d.add_argument("--section-selector", help="CSS selector for section headings, if sections are not detected")
    d.add_argument("--prefix", help="only keep links whose path starts with this, e.g. /linux-unix/")
    d.add_argument("--max-pages", type=int, default=200)

    fmt = argparse.ArgumentParser(add_help=False)
    f = fmt.add_argument_group("what to include")
    f.add_argument("--sections", metavar="SPEC", help='sections to keep by number: "1-5,8" or "not 12,13"')
    f.add_argument("--skip-sections", metavar="REGEX", help="drop sections whose name matches REGEX")
    f.add_argument("--select", metavar="SPEC", help="only these topic numbers, e.g. 1-5,8")
    f.add_argument("--include", metavar="REGEX", help="keep topics whose address or title matches")
    f.add_argument("--exclude", metavar="REGEX", help="drop topics whose address or title matches, e.g. 'quiz|interview'")
    o = fmt.add_argument_group("document")
    o.add_argument("--title"); o.add_argument("--subtitle"); o.add_argument("--author")
    o.add_argument("--out-dir", type=Path, default=Path("output"))
    o.add_argument("--page-size", choices=["A4", "Letter"], default="A4")
    o.add_argument("--landscape", action="store_true")
    o.add_argument("--font-size", type=float, default=10.5, help="body font size in pt")
    o.add_argument("--margin-mm", type=float, default=20.0)
    o.add_argument("--code-cols", type=int, default=90, help="wrap code lines longer than this (with a visible marker)")
    o.add_argument("--toc-depth", type=int, choices=[1, 2], help="1 = chapters only, 2 = also headings (default: auto)")
    o.add_argument("--no-chapter-breaks", action="store_true", help="do not start every chapter on a new page")
    o.add_argument("--no-images", action="store_true"); o.add_argument("--no-cover", action="store_true")
    o.add_argument("--no-toc", action="store_true"); o.add_argument("--no-source-urls", action="store_true")
    o.add_argument("--no-report-page", action="store_true", help="omit the build report appendix from the PDF")
    o.add_argument("--docx", action="store_true", help="also write an editable Word document (.docx)")
    o.add_argument("--no-pdf", action="store_true", help="skip the PDF (use with --docx for a Word-only build)")
    o.add_argument("--markdown", action="store_true", help="also write an editable Markdown file")
    o.add_argument("--html-export", action="store_true", help="also write the HTML used for the PDF")
    o.add_argument("--allow-partial", action="store_true", help="build even if some chapters failed (labelled PARTIAL)")
    o.add_argument("-y", "--yes", action="store_true", help="do not ask questions; use sensible defaults")

    ap = argparse.ArgumentParser(prog="tsf.py", description=__doc__.split("\n\n")[0], epilog=EXAMPLES,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--version", action="version", version=f"{APP_NAME} (TSF) {__version__}")
    sub = ap.add_subparsers(dest="cmd", required=True, metavar="command")

    gp = sub.add_parser("go", parents=[common, disc, fmt], help="guided mode (default): paste an address, answer a few questions")
    gp.add_argument("go_url", nargs="?", metavar="URL", help="address of the tutorial's main page")

    dp = sub.add_parser("discover", parents=[common, disc], help="list the topics on a tutorial page into an editable file")
    dp.add_argument("url")
    dp.add_argument("-o", "--output", type=Path, default=Path("topics.txt"))

    b = sub.add_parser("build", parents=[common, disc, fmt], help="build a PDF from an address, a topic list, or saved pages")
    src = b.add_mutually_exclusive_group(required=True)
    src.add_argument("--index", metavar="URL", help="tutorial main page: find its topics and build all of them")
    src.add_argument("--chapters", type=Path, metavar="FILE", help="topic list written by 'discover' or guided mode")
    src.add_argument("--url", action="append", metavar="URL", help="article address (repeatable)")
    src.add_argument("--html", nargs="+", metavar="PATH", help="HTML files or folders you saved yourself")
    b.add_argument("--base-url", help="with --html: base address used to resolve relative links")

    i = sub.add_parser("inspect", parents=[common], help="show how one page (address or saved file) is read")
    i.add_argument("target")

    sub.add_parser("sites", help="list the sites with built-in rules and example start addresses")
    st = sub.add_parser("selftest", help="build a small sample document to verify your installation")
    st.add_argument("--out-dir", type=Path, default=Path("output"))
    return ap


def settings_from(a: argparse.Namespace) -> Settings:
    s = Settings()
    g = lambda k, d=None: getattr(a, k, d)
    s.delay, s.timeout = g("delay", 1.0), g("timeout", 30.0)
    if g("cache_dir"):
        s.cache_dir = g("cache_dir")
    s.refresh, s.offline = bool(g("refresh")), bool(g("offline"))
    s.ack_terms, s.allow_private = [fix_host(x) for x in (g("ack_terms", []) or [])], bool(g("allow_private"))
    s.profiles_file, s.selector, s.drop_text = g("profiles"), g("selector"), g("drop_text", []) or []
    s.section_selector = g("section_selector")
    s.min_words, s.verbose = g("min_words", 30), bool(g("verbose"))
    s.max_pages = g("max_pages", 200)
    s.out_dir = g("out_dir", Path("output")) or Path("output")
    s.title, s.subtitle, s.author = g("title") or "", g("subtitle") or "", g("author") or ""
    s.page_size, s.landscape = g("page_size", "A4"), bool(g("landscape"))
    s.font_pt, s.margin_mm = g("font_size", 10.5), g("margin_mm", 20.0)
    s.code_cols = max(40, g("code_cols", 90))
    s.toc_depth, s.chapter_breaks = g("toc_depth"), not g("no_chapter_breaks", False)
    s.images = not g("no_images", False)
    s.cover, s.toc = not g("no_cover", False), not g("no_toc", False)
    s.source_urls, s.report_page = not g("no_source_urls", False), not g("no_report_page", False)
    s.write_markdown, s.write_html = bool(g("markdown")), bool(g("html_export"))
    s.docx, s.pdf = bool(g("docx")), not g("no_pdf", False)
    s.allow_partial, s.yes = bool(g("allow_partial")), bool(g("yes"))
    return s


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        argv = ["go"]
    elif argv[0] not in COMMANDS | {"-h", "--help", "--version"}:
        argv = ["go"] + argv                    # bare address (or bare options) -> guided mode
    a = make_parser().parse_args(argv)
    s = settings_from(a)
    logging.basicConfig(level=logging.DEBUG if s.verbose else logging.WARNING, format="%(levelname)s %(message)s")
    if not s.verbose:
        for noisy in ("weasyprint", "fontTools", "fontTools.subset"):
            logging.getLogger(noisy).setLevel(logging.ERROR)
    try:
        if a.cmd == "sites":
            return cmd_sites()
        if a.cmd == "selftest":
            return cmd_selftest(s)
        if not s.pdf and not s.docx:
            say("Nothing to build: --no-pdf needs --docx.")
            return 2
        fetcher = Fetcher(s)
        if a.cmd == "go":
            return run_guided(a, s, fetcher)
        if a.cmd == "inspect":
            tgt = a.target if Path(a.target).exists() else fix_url(a.target)
            return cmd_inspect(tgt, s, fetcher)
        if a.cmd == "discover":
            url = fix_url(a.url)
            host = urllib.parse.urlsplit(url).hostname or ""
            if not host_acknowledged(s, host):
                return 1
            chapters, info = discover(fetcher, url, s, a.link_selector, a.prefix)
            if not chapters:
                say("No topic links found. Try --link-selector or 'inspect', or list article addresses yourself with build --url.")
                return 1
            write_chapters_file(a.output, chapters, info)
            print_sections(chapters, info, host)
            say(f"\nWrote {a.output}. Edit it (delete or reorder lines), then run:\n"
                f"  python3 tsf.py build --chapters {a.output}")
            return 0
        # ---- build ----
        local, index_url = False, None
        if a.index:
            url = fix_url(a.index)
            host = urllib.parse.urlsplit(url).hostname or ""
            if not host_acknowledged(s, host):
                return 1
            chapters, info = discover(fetcher, url, s, a.link_selector, a.prefix)
            index_url = info["index_url"]
            say(f"Discovered {len(chapters)} topics ({info['method']}, confidence {info['confidence']}).")
            if not chapters:
                return 1
            s.title = s.title or info.get("page_title", "")
            s.subtitle = s.subtitle or info.get("page_subtitle", "")
        elif a.chapters:
            chapters = read_chapters_file(a.chapters)
            meta = read_chapters_file.meta
            s.title = s.title or meta.get("title", "")
            s.subtitle = s.subtitle or meta.get("subtitle", "")
        elif a.url:
            chapters = [Chapter(i, fix_url(u), "", 0, "cli") for i, u in enumerate(a.url, 1)]
        else:
            chapters, local = chapters_from_html_paths(a.html), True
        chapters = filter_sections(chapters, a.sections, a.skip_sections)
        chapters = apply_filters(chapters, a, s)
        return build_pipeline(chapters, s, None if local else fetcher, local_mode=local,
                              base_url=a.base_url if local else None, source_index_url=index_url)
    except FetchError as e:
        say(f"Error: {e.message}\n  {e.hint}".rstrip())
        return 1
    except UrlError as e:
        say(f"Error: {e}")
        return 2
    except KeyboardInterrupt:
        say("\nInterrupted. Pages already downloaded are cached; run the same command again to continue.")
        return 130
    except ValueError as e:
        say(f"Error: {e}")
        return 2


if __name__ == "__main__":
    sys.exit(main())
