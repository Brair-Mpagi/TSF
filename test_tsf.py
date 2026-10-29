"""Offline tests: unit tests plus an end-to-end run against a local fixture site.

Run:  pip install pytest && pytest -q
"""
import json
import socket
import struct
import sys
import threading
import zlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import tsf as t  # noqa: E402


# ---------------------------------------------------------------- unit tests
def test_normalize_and_dedupe():
    assert t.normalize_url("HTTPS://Example.com:443/a/b?utm_source=x&id=3#frag") == "https://example.com/a/b?id=3"
    assert t.dedupe_key("https://www.x.com/a/") == t.dedupe_key("http://x.com/a")
    assert t.dedupe_key("https://x.com/a/index.html") == t.dedupe_key("https://x.com/a")


@pytest.mark.parametrize("url", ["ftp://x.com/a", "file:///etc/passwd", "javascript:alert(1)", "http://user:pw@x.com/",
                                 "https://x.com:8080/", "http:///nohost"])
def test_bad_urls_rejected(url):
    with pytest.raises(t.UrlError):
        t.validate_url_syntax(url)


@pytest.mark.parametrize("host", ["127.0.0.1", "localhost", "10.0.0.5", "192.168.1.1", "169.254.169.254", "::1",
                                  "0.0.0.0", "100.64.0.1", "[::ffff:127.0.0.1]"])
def test_private_hosts_blocked(host):
    with pytest.raises(t.FetchError) as e:
        t.assert_public_host(host.strip("[]"), 80)
    assert e.value.kind in ("ssrf", "dns")


def test_code_text_is_exact():
    from bs4 import BeautifulSoup
    html = ('<pre><span class="linenos">1\n2</span><code>def f(x):\n    if x:\n\t\treturn "a  b"\n'
            '    return None<br>  last</code></pre>')
    el = BeautifulSoup(html, "lxml").pre
    assert t.code_text(el) == 'def f(x):\n    if x:\n\t\treturn "a  b"\n    return None\n  last'


def test_code_text_div_lines_and_nbsp():
    from bs4 import BeautifulSoup
    el = BeautifulSoup('<pre><div class="line">a = 1</div><div class="line">&nbsp;&nbsp;b = 2</div></pre>', "lxml").pre
    assert t.code_text(el) == "a = 1\n  b = 2"


def test_wrap_code_marks_continuation_and_reports_count():
    long = "    " + "x = call(" + ", ".join(f"arg{i}" for i in range(30)) + ")"
    out, n = t.wrap_code(long, 60)
    lines = out.split("\n")
    assert n >= 1 and all(len(l) <= 60 for l in lines)
    assert lines[1].lstrip().startswith(t.CONT_MARK)
    assert out.replace("\n", "").replace(t.CONT_MARK + " ", "").replace(" ", "") == long.replace(" ", "")


def test_selection_parser():
    assert t.parse_selection("1-3,5", 6) == {1, 2, 3, 5}
    with pytest.raises(ValueError):
        t.parse_selection("9", 3)


def test_chapters_file_roundtrip(tmp_path):
    chs = [t.Chapter(1, "https://a.com/x", "One", 0), t.Chapter(2, "https://a.com/y", "Two", 1)]
    f = tmp_path / "c.txt"
    t.write_chapters_file(f, chs, {"index_url": "https://a.com/", "method": "m", "confidence": "high", "notes": []})
    back = t.read_chapters_file(f)
    assert [(c.url, c.title, c.depth) for c in back] == [(c.url, c.title, c.depth) for c in chs]


def test_safe_join_blocks_traversal(tmp_path):
    with pytest.raises(ValueError):
        t.safe_join(tmp_path, "..", "evil")


def test_junk_removed_but_code_in_widget_kept():
    s = t.Settings()
    page = b"""<html><body><article><h1>T</h1><p>%s</p>
    <div class="related-posts"><a href="/a">A</a> <a href="/b">B</a></div>
    <div class="sidebar-note"><pre>keep_this()</pre></div></article></body></html>""" % (b"word " * 60)
    art = t.extract_article(page, "https://x.com/p", "https://x.com/p", s)
    text = t.plain_text(art["blocks"])
    assert "keep_this()" in text and "A B" not in text


def test_access_denied_page_is_not_an_article():
    page = b"<html><body><main><h1>Just a moment...</h1><p>Verify you are human to continue.</p></main></body></html>"
    art = t.extract_article(page, "https://x.com/p", "https://x.com/p", t.Settings())
    assert art["status"] == "failed" and art["error"]["kind"] in ("access_page", "empty")


# ---------------------------------------------------------------- local fixture site
def _png() -> bytes:
    w = h = 40
    raw = b"".join(b"\x00" + bytes((x * 6 + y) % 256 for x in range(w) for _ in range(3)) for y in range(h))
    def chunk(k, d):
        return struct.pack(">I", len(d)) + k + d + struct.pack(">I", zlib.crc32(k + d) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + \
        chunk(b"IDAT", zlib.compress(raw, 0)) + chunk(b"IEND", b"")


LOREM = ("Relational databases store data in tables and queries describe the result you want rather than the steps "
         "to compute it. ") * 4


def article(title: str, body: str) -> bytes:
    return (f"<html><head><title>{title} | Site</title><link rel='canonical' href='/sql/{title.lower()}'></head><body>"
            f"<nav><a href='/'>Home</a></nav><div class='ad-slot'>BUY NOW <a href='/ad'>ad</a></div>"
            f"<article><h1>{title}</h1><p>{LOREM}</p>{body}</article>"
            f"<div class='share'>SHARE THIS PAGE <a href='#'>x</a></div></body></html>").encode()


PAGES = {
    "/sql/": None,  # index built below
    "/sql/intro": article("Intro", "<pre class='language-sql'>SELECT *\n  FROM t\n WHERE x = 1;</pre>"
                                   "<img src='/img/a.png' alt='diagram' width='40' height='40'>"),
    "/sql/select": article("Select", "<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table>"
                                     "<pre>def f():\n    return 1</pre>"),
    "/sql/joins": article("Joins", "<pre>JOIN a ON a.id = b.id</pre>"),
    "/sql/dupe": article("Intro", "<pre class='language-sql'>SELECT *\n  FROM t\n WHERE x = 1;</pre>"
                                  "<img src='/img/a.png' alt='diagram' width='40' height='40'>"),
    "/sql/captcha": b"<html><body><main><h1>Just a moment...</h1><p>Verify you are human.</p></main></body></html>",
}
INDEX = ("<html><head><title>SQL Tutorial</title></head><body><header><a href='/login'>Login</a></header><main>"
         "<h1>SQL Tutorial</h1><ul>" +
         "".join(f"<li><a href='{u}'>{n}</a></li>" for u, n in [
             ("/sql/intro", "Intro"), ("/sql/select", "Select"), ("/sql/intro/", "Intro again"), ("/sql/joins", "Joins"),
             ("/sql/missing", "Missing chapter"), ("/sql/blocked", "Blocked chapter"), ("/private/secret", "Robots chapter"),
             ("/sql/captcha", "Captcha chapter"), ("/sql/dupe", "Duplicate content")]) +
         "</ul></main><footer><a href='/terms'>Terms</a></footer></body></html>").encode()


GUIDE = ("<html><head><title>Guide Tutorial</title></head><body><header><a href='/login'>Login</a></header><main>"
         "<h1>Guide Tutorial</h1><p>Last Updated : 11 Aug, 2026</p>"
         "<p>This guide teaches the essentials quickly and clearly for learners everywhere.</p>"
         "<h2>Basics</h2><p>Core concepts that build a strong foundation for everything else.</p>"
         "<ul><li><a href='/sql/intro'>Introduction</a></li><li><a href='/sql/select'>Select</a></li></ul>"
         "<h2>Advanced</h2><p>Deeper features for experienced users who need more power.</p>"
         "<ul><li><a href='/sql/joins'>Joins</a></li></ul>"
         "<h2>Related Articles</h2><ul><li><a href='/sql/dupe'>Another page</a></li></ul>"
         "</main></body></html>").encode()


W3_LINKS = [("default.asp", "SQL HOME"), ("sql_intro.asp", "SQL Intro"), ("sql_syntax.asp", "SQL Syntax"),
            ("sql_select.asp", "SQL SELECT"), ("sql_where.asp", "SQL WHERE"), ("sql_create_db.asp", "SQL Create DB"),
            ("sql_drop_db.asp", "SQL Drop DB"), ("sql_create_table.asp", "SQL Create Table")]


def w3_page(slug: str, title: str) -> bytes:
    menu = ("<div id='mySidenav'><div id='leftmenuinner'><h2 class='left'>SQL Tutorial</h2>" +
            "".join(f"<a href='{u}'>{n}</a>" for u, n in W3_LINKS[:5]) +
            "<h2 class='left'>SQL Database</h2>" + "".join(f"<a href='{u}'>{n}</a>" for u, n in W3_LINKS[5:]) +
            "</div></div><div id='nav_tutorials'>" + "".join(f"<a href='/other{i}/x.asp'>Other {i}</a>" for i in range(30)) +
            "</div>")
    return (f"<html><head><title>{title}</title></head><body>{menu}<div id='mainLeaderboard'>AD AD AD</div>"
            f"<div id='main'><h1>{title}</h1><p>{LOREM}</p>"
            f"<div class='w3-example'><h3>Example</h3><div class='w3-code notranslate'>SELECT *<br>&nbsp;&nbsp;FROM Customers<br>"
            f"&nbsp;&nbsp;WHERE Country = 'UK';</div><a class='w3-btn' href='/tryit.asp?x=1'>Try it Yourself &raquo;</a></div>"
            f"<div class='nextprev'><a href='sql_intro.asp'>&laquo; Previous</a><a href='sql_syntax.asp'>Next &raquo;</a></div>"
            f"</div></body></html>").encode()


class Handler(BaseHTTPRequestHandler):
    hits: dict = {}

    def log_message(self, *a):
        pass

    def do_GET(self):
        Handler.hits[self.path] = Handler.hits.get(self.path, 0) + 1
        path = self.path.split("?")[0]
        if path == "/robots.txt":
            return self._send(200, b"User-agent: *\nDisallow: /private/\n", "text/plain")
        if path == "/img/a.png":
            return self._send(200, _png(), "image/png")
        if path == "/sql/":
            return self._send(200, INDEX)
        if path == "/guide/":
            return self._send(200, GUIDE)
        if path.startswith("/w3/sql/") and path.endswith(".asp"):
            slug = path.rsplit("/", 1)[1]
            names = dict(W3_LINKS)
            if slug in names:
                return self._send(200, w3_page(slug, names[slug]))
        if path == "/sql/blocked":
            return self._send(403, b"forbidden")
        if path in PAGES and PAGES[path]:
            return self._send(200, PAGES[path])
        self._send(404, b"nope")

    def _send(self, code, body, ctype="text/html; charset=utf-8"):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


@pytest.fixture()
def site(monkeypatch):
    monkeypatch.setenv("NO_PROXY", "127.0.0.1,localhost")
    monkeypatch.setenv("no_proxy", "127.0.0.1,localhost")
    Handler.hits = {}
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()


def settings(tmp_path, **kw):
    s = t.Settings(out_dir=tmp_path / "out", cache_dir=tmp_path / "cache", config_dir=tmp_path / "cfg",
                   delay=0, allow_private=True, retries=0)
    for k, v in kw.items():
        setattr(s, k, v)
    return s


def test_ssrf_blocks_loopback_by_default(tmp_path, site):
    s = settings(tmp_path, allow_private=False)
    with pytest.raises(t.FetchError) as e:
        t.Fetcher(s).get(site + "/sql/")
    assert e.value.kind in ("url", "ssrf")      # port not allowed, or address not public


def test_discover_finds_chapters_and_dedupes(tmp_path, site):
    s = settings(tmp_path)
    chapters, info = t.discover(t.Fetcher(s), site + "/sql/", s)
    urls = [c.url.replace(site, "") for c in chapters]
    assert urls[:3] == ["/sql/intro", "/sql/select", "/sql/joins"]
    assert "/login" not in urls and "/terms" not in urls
    assert len(urls) == len(set(urls)) == 8       # "/sql/intro/" collapsed into "/sql/intro"
    assert any("duplicate" in n for n in info["notes"])


def test_full_build_reports_failures_honestly(tmp_path, site):
    s = settings(tmp_path, title="Fixture Guide", write_markdown=True)
    f = t.Fetcher(s)
    chapters, _ = t.discover(f, site + "/sql/", s)
    # partial not confirmed -> exit code 3, no PDF
    rc = t.build_pipeline(chapters, s, f)
    out = tmp_path / "out" / "fixture-guide"
    assert rc == 3 and not list(out.glob("*.pdf"))
    comp = json.loads((out / "completeness.json").read_text())
    by = {c["title"]: c for c in comp["chapters"]}
    assert by["Missing chapter"]["status"] == "failed" and "404" in by["Missing chapter"]["notes"][0]
    assert by["Blocked chapter"]["status"] == "failed"
    assert by["Robots chapter"]["status"] == "failed" and "robots" in by["Robots chapter"]["notes"][0]
    assert by["Captcha chapter"]["status"] == "failed"
    assert by["Duplicate content"]["status"] == "skipped"
    assert "/private/secret" not in Handler.hits           # robots.txt was honoured: never requested
    # now confirm partial
    s.allow_partial = True
    rc = t.build_pipeline(chapters, s, f)
    pdf = out / "Fixture-Guide-PARTIAL.pdf"
    assert rc == 0 and pdf.exists()
    manifest = json.loads((out / "manifest.json").read_text())
    assert manifest["status"] == "Partial"
    md = (out / "Fixture-Guide.md").read_text()
    assert "SELECT *\n  FROM t\n WHERE x = 1;" in md      # code byte-exact
    assert "BUY NOW" not in md and "SHARE THIS PAGE" not in md
    assert (out / "assets").exists() and any((out / "assets").iterdir())   # image downloaded
    from pypdf import PdfReader
    text = "".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
    assert "Missing chapter" in text and "missing from this document" in text


def test_cache_prevents_refetch(tmp_path, site):
    s = settings(tmp_path)
    f = t.Fetcher(s)
    f.get(site + "/sql/intro")
    f2 = t.Fetcher(s)
    r = f2.get(site + "/sql/intro")
    assert r.from_cache and Handler.hits["/sql/intro"] == 1


def test_403_is_not_retried_and_host_stops_after_repeats(tmp_path, site):
    s = settings(tmp_path, retries=3)
    f = t.Fetcher(s)
    for _ in range(3):
        with pytest.raises(t.FetchError) as e:
            f.get(site + "/sql/blocked")
        assert e.value.kind == "access_denied"
    assert Handler.hits["/sql/blocked"] == 3                # exactly one attempt each, no retries
    with pytest.raises(t.FetchError) as e:
        f.get(site + "/sql/intro")
    assert e.value.kind == "host_blocked"


def test_import_html_dir(tmp_path):
    d = tmp_path / "html"
    d.mkdir()
    (d / "10-b.html").write_bytes(article("Beta", "<pre>b()</pre>"))
    (d / "2-a.html").write_bytes(article("Alpha", "<pre>a()</pre>"))
    chs = t.chapters_from_html_paths([str(d)])
    assert [c.path.name for c in chs] == ["2-a.html", "10-b.html"]      # natural order
    s = settings(tmp_path, title="Imported")
    assert t.build_pipeline(chs, s, None, local_mode=True) == 0
    comp = json.loads((tmp_path / "out" / "imported" / "completeness.json").read_text())
    assert comp["status"] in ("Completed", "Completed with warnings") and comp["summary"]["included"] == 2


def test_selftest_builds_and_verifies(tmp_path):
    s = settings(tmp_path)
    assert t.cmd_selftest(s) == 0
    comp = json.loads((tmp_path / "out" / "selftest-sample-guide" / "completeness.json").read_text())
    assert all(c["ok"] for c in comp["verification"])


# ---------------------------------------------------------------- sections / parts / guided mode
def test_fix_url_accepts_messy_input_and_rejects_placeholders():
    assert t.fix_url("www.geeksforgeeks.org/sql-tutorial/") == "https://www.geeksforgeeks.org/sql-tutorial/"
    assert t.fix_url("[x](https://a.com/p)") == "https://a.com/p"
    assert t.fix_url("<https://a.com/p>") == "https://a.com/p"
    with pytest.raises(t.UrlError) as e:
        t.fix_url("LINUX_URL")
    assert "copy the address" in str(e.value)
    assert t.fix_host("https://www.GeeksforGeeks.org/x") == "www.geeksforgeeks.org"


def test_discover_groups_topics_into_sections(tmp_path, site):
    s = settings(tmp_path)
    chapters, info = t.discover(t.Fetcher(s), site + "/guide/", s)
    assert t.sections_of(chapters) == ["Basics", "Advanced"]          # "Related Articles" dropped
    assert [(c.title, c.section) for c in chapters] == [("Introduction", "Basics"), ("Select", "Basics"),
                                                         ("Joins", "Advanced")]
    assert chapters[0].section_desc.startswith("Core concepts")
    assert info["page_title"] == "Guide Tutorial" and info["page_subtitle"].startswith("This guide teaches")
    assert any("Related Articles" in n for n in info["notes"])


def test_topic_file_roundtrip_keeps_sections(tmp_path, site):
    s = settings(tmp_path)
    chapters, info = t.discover(t.Fetcher(s), site + "/guide/", s)
    f = tmp_path / "topics.txt"
    t.write_chapters_file(f, chapters, info)
    assert "== Basics" in f.read_text()
    back = t.read_chapters_file(f)
    assert [(c.title, c.section, c.section_desc) for c in back] == [(c.title, c.section, c.section_desc) for c in chapters]
    assert t.read_chapters_file.meta["title"] == "Guide Tutorial"
    assert t.read_chapters_file.meta["subtitle"].startswith("This guide teaches")


def test_filter_sections():
    chs = [t.Chapter(1, "u1", "a", section="A"), t.Chapter(2, "u2", "b", section="B"), t.Chapter(3, "u3", "c", section="C")]
    assert [c.title for c in t.filter_sections(chs, "1,3")] == ["a", "c"]
    assert [c.title for c in t.filter_sections(chs, "not 2")] == ["a", "c"]
    assert [c.title for c in t.filter_sections(chs, None, "^B$")] == ["a", "c"]
    with pytest.raises(ValueError):
        t.filter_sections(chs, "9")


def test_guided_mode_builds_parts_pdf_and_markdown(tmp_path, site):
    out = tmp_path / "out"
    rc = t.main(["go", site + "/guide/", "--yes", "--markdown", "--allow-private", "--delay", "0",
                 "--cache-dir", str(tmp_path / "cache"), "--out-dir", str(out), "--allow-partial"])
    assert rc == 0
    d = out / "guide-tutorial"
    pdf = d / "Guide-Tutorial.pdf"
    assert pdf.exists() and (d / "topics.txt").exists()
    from pypdf import PdfReader
    rd = PdfReader(str(pdf))
    text = "".join(p.extract_text() for p in rd.pages)
    assert "Guide Tutorial" in text and "This guide teaches the essentials" in text      # cover title + subtitle
    assert "PART 1" in text.upper() and "Basics" in text and "Core concepts that build" in text
    toc_page = rd.pages[1].extract_text()
    assert "Contents" in toc_page and "Basics" in toc_page and "Advanced" in toc_page and "Joins" in toc_page  # one page
    top = [o.title for o in rd.outline if not isinstance(o, list)]
    assert "Basics" in top and "Advanced" in top and "Intro" not in top                  # chapters nest under parts
    md = (d / "Guide-Tutorial.md").read_text()
    assert md.index("## Basics") < md.index("### Intro") < md.index("## Advanced") < md.index("### Joins")
    comp = json.loads((d / "completeness.json").read_text())
    assert comp["summary"]["included"] == 3 and comp["chapters"][0]["section"] == "Basics"


def test_markdown_code_comments_are_not_turned_into_headings():
    blocks = [{"type": "heading", "level": 2, "html": "Sub"},
              {"type": "code", "language": "python", "text": "# a comment\nprint(1)"}]
    md = "\n".join(t.md_blocks(blocks, hshift=2))
    assert "#### Sub" in md and "\n# a comment\n" in md and "## a comment" not in md


def test_sections_flag_selects_by_number(tmp_path, site):
    out = tmp_path / "out"
    rc = t.main(["build", "--index", site + "/guide/", "--sections", "2", "--allow-private", "--delay", "0",
                 "--cache-dir", str(tmp_path / "cache"), "--out-dir", str(out), "--allow-partial", "--ack-terms", "127.0.0.1"])
    assert rc == 0
    comp = json.loads((out / "guide-tutorial" / "completeness.json").read_text())
    assert [c["title"] for c in comp["chapters"]] == ["Joins"]


def test_bare_placeholder_gives_friendly_error(capsys):
    rc = t.main(["discover", "LINUX_URL"])
    assert rc == 2 and "not a web address" in capsys.readouterr().err


# ---------------------------------------------------------------- Word (.docx) output
from docx import Document as _Doc  # noqa: E402


def _synthetic_articles(blocks, title="Chapter A", section="", desc=""):
    return [{"schema": 1, "index": 1, "url": "https://example.org/a", "canonical_url": None, "final_url": None,
             "title": title, "section": section, "section_desc": desc, "source_name": "x", "depth": 0, "status": "ok",
             "confidence": 1.0, "method": "", "fetched_at": None, "from_cache": False, "content_hash": "h", "words": 99,
             "counts": t.count_blocks(blocks), "warnings": [], "error": None, "blocks": blocks}]


def _docx_for(tmp_path, blocks, **kw):
    s = settings(tmp_path, **kw)
    arts = _synthetic_articles(blocks)
    out = tmp_path / "x.docx"
    meta = {"title": "Doc", "date": "2026-01-01"}
    t.build_docx(arts, s, meta, "Completed", t.report_rows(arts), tmp_path / "assets", out)
    return out, arts, s, meta


def P(html):
    return {"type": "paragraph", "html": html}


def test_docx_code_is_exact_including_tabs_and_leading_spaces(tmp_path):
    code = "def f(x):\n    if x:\n\t\treturn  'a   b'\n  # comment\n\n    return None  "
    out, arts, s, meta = _docx_for(tmp_path, [P("intro " * 10), {"type": "code", "language": "python", "text": code}])
    d = _Doc(str(out))
    assert [p.text for p in d.paragraphs if p.style.name == "Code"] == [code]
    assert all(c["ok"] for c in t.verify_docx(out, arts, meta, s))


def test_docx_ordered_lists_restart_and_nest(tmp_path):
    lst = lambda items, ordered=True: {"type": "list", "ordered": ordered, "start": 1, "items": items}
    blocks = [lst([[P("one")], [P("two"), lst([[P("two-a")]], False)]]), P("between"), lst([[P("again one")]])]
    out, *_ = _docx_for(tmp_path, blocks)
    d = _Doc(str(out))
    items = [p for p in d.paragraphs if p.style.name == "List Paragraph"][:4]      # the Sources appendix follows
    ids = [p._p.pPr.numPr.numId.val for p in items]
    assert [p.text for p in items] == ["one", "two", "two-a", "again one"]
    assert ids[0] == ids[1] and ids[3] != ids[0]                     # second ordered list gets its own numbering
    assert items[2]._p.pPr.numPr.ilvl.val == 1
    numbering = d.part.numbering_part.element.xml
    assert numbering.count("startOverride") >= 2


def test_docx_table_merged_cells_and_header(tmp_path):
    cell = lambda h, th=False, cs=1, rs=1: {"html": h, "th": th, "colspan": cs, "rowspan": rs}
    tb = {"type": "table", "caption": "Joins", "rows": [
        {"head": True, "cells": [cell("Join", True), cell("Behaviour", True, cs=2)]},
        {"head": False, "cells": [cell("INNER", rs=2), cell("a"), cell("b")]},
        {"head": False, "cells": [cell("c"), cell("<strong>d</strong>")]}]}
    out, *_ = _docx_for(tmp_path, [tb])
    d = _Doc(str(out))
    tbl = d.tables[0]
    assert len(tbl.rows) == 3 and len(tbl.columns) == 3
    assert tbl.cell(1, 0).text == "INNER" and tbl.cell(2, 2).text == "d"
    assert "tblHeader" in tbl._tbl.xml and "Joins" in [p.text for p in d.paragraphs if p.style.name == "Caption"]


def test_docx_inline_formatting_links_and_invalid_chars(tmp_path):
    html = 'Use <code>SELECT</code> and <strong>bold</strong> <em>it</em> <a href="https://e.org/x">link</a> x<sup>2</sup>\x00\x0b'
    out, *_ = _docx_for(tmp_path, [P(html)])
    d = _Doc(str(out))
    p = [x for x in d.paragraphs if x.text.startswith("Use")][0]
    assert p.runs[1].style.name == "Inline Code" and p.runs[3].bold
    assert "https://e.org/x" in [r.target_ref for r in d.part.rels.values() if r.is_external]
    assert "\x00" not in p.text


def test_docx_images_embedded_and_missing_marked(tmp_path):
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "a.png").write_bytes(_png())
    blocks = [{"type": "image", "src": "a.png", "url": None, "alt": "d", "caption": "Figure 1", "asset": "a.png"},
              {"type": "image", "src": "b.png", "url": None, "alt": "gone", "caption": "", "asset": None}]
    out, arts, s, meta = _docx_for(tmp_path, blocks)
    d = _Doc(str(out))
    assert len(d.inline_shapes) == 1
    assert any("Image not available" in p.text for p in d.paragraphs)
    assert all(c["ok"] for c in t.verify_docx(out, arts, meta, s))


def test_docx_parts_and_contents_field(tmp_path):
    s = settings(tmp_path)
    arts = _synthetic_articles([P("text " * 20)], section="Basics", desc="Core ideas.")
    out = tmp_path / "p.docx"
    t.build_docx(arts, s, {"title": "Doc", "date": "2026-01-01"}, "Completed", t.report_rows(arts), tmp_path / "a", out)
    d = _Doc(str(out))
    kinds = [(p.style.name, p.text) for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert ("Heading 1", "Basics") in kinds and ("Heading 2", "Chapter A") in kinds
    assert "Core ideas." in [p.text for p in d.paragraphs if p.style.name == "Part Description"]
    assert "TOC \\o" in d.element.xml and "updateFields" in d.settings.element.xml


def test_cli_docx_only_build_and_flag_validation(tmp_path, site, capsys):
    out = tmp_path / "out"
    base = ["--allow-private", "--delay", "0", "--cache-dir", str(tmp_path / "cache"), "--out-dir", str(out), "--allow-partial"]
    assert t.main(["go", site + "/guide/", "--yes", "--docx", "--no-pdf"] + base) == 0
    d = out / "guide-tutorial"
    assert (d / "Guide-Tutorial.docx").exists() and not list(d.glob("*.pdf"))
    manifest = json.loads((d / "manifest.json").read_text())
    assert manifest["files"] == {"pdf": None, "docx": "Guide-Tutorial.docx"}
    doc = _Doc(str(d / "Guide-Tutorial.docx"))
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert ("Heading 1", "Basics") in heads and ("Heading 2", "Intro") in heads and ("Heading 1", "Advanced") in heads
    assert t.main(["go", site + "/guide/", "--yes", "--no-pdf"] + base) == 2
    assert "needs --docx" in capsys.readouterr().err


# ---------------------------------------------------------------- other sites: side menus, profiles
def _w3_profile(tmp_path):
    f = tmp_path / "profiles.json"
    f.write_text(json.dumps({"127.0.0.1": {
        "content": ["#main"], "toc": ["#leftmenuinner"],
        "remove": ["a.w3-btn", ".nextprev", "#mainLeaderboard"], "code_selectors": ["div.w3-code"]}}))
    return f


def test_discover_side_menu_without_any_profile(tmp_path, site):
    s = settings(tmp_path)
    chapters, info = t.discover(t.Fetcher(s), site + "/w3/sql/sql_intro.asp", s)
    assert info["method"] == "sidebar menu"
    assert [c.title for c in chapters] == [n for _, n in W3_LINKS]          # includes the page itself and SQL HOME
    assert not any("other" in c.url for c in chapters)                      # the site-wide mega menu is ignored
    assert t.sections_of(chapters) == ["SQL Tutorial", "SQL Database"]


def test_discover_with_profile_and_code_selector_extraction(tmp_path, site):
    s = settings(tmp_path, profiles_file=_w3_profile(tmp_path), title="W3 Like")
    f = t.Fetcher(s)
    chapters, info = t.discover(f, site + "/w3/sql/", s) if False else t.discover(f, site + "/w3/sql/sql_syntax.asp", s)
    assert info["method"].startswith("profile") and len(chapters) == len(W3_LINKS)
    res = f.get(site + "/w3/sql/sql_select.asp")
    art = t.extract_article(res.body, res.final_url, res.final_url, s, profile_host="127.0.0.1")
    code = [b for b in art["blocks"] if b["type"] == "code"]
    assert code and code[0]["text"] == "SELECT *\n  FROM Customers\n  WHERE Country = 'UK';"
    text = t.plain_text(art["blocks"])
    assert "Try it Yourself" not in text and "AD AD AD" not in text and "Previous" not in text


def test_builtin_profiles_cover_major_sites_and_sites_command(capsys):
    for host in ("geeksforgeeks.org", "w3schools.com", "tutorialspoint.com", "javatpoint.com", "programiz.com",
                 "developer.mozilla.org", "docs.python.org", "realpython.com", "baeldung.com", "docs.oracle.com"):
        prof = t.load_profile("www." + host, t.Settings())
        assert prof["name"] != "any other site" and prof["content"], host
    assert t.load_profile("foo.readthedocs.io", t.Settings())["toc"]
    assert t.load_profile("www.w3schools.com", t.Settings())["code_selectors"] == ["div.w3-code"]
    assert "a.headerlink" in t.load_profile("example.org", t.Settings())["remove"]       # generic rules always apply
    assert t.main(["sites"]) == 0
    out = capsys.readouterr().err
    assert "W3Schools" in out and "JavaTpoint" in out and "TutorialsPoint" in out


def test_chapter_page_pasted_gives_tutorial_title_not_chapter_title(tmp_path, site):
    s = settings(tmp_path)
    chapters, info = t.discover(t.Fetcher(s), site + "/w3/sql/sql_where.asp", s)
    assert info["page_title"] == "SQL Tutorial" and info["page_subtitle"] == ""


def test_duplicates_are_reported_but_not_in_the_documents(tmp_path, site):
    out = tmp_path / "out"
    base = ["--allow-private", "--delay", "0", "--cache-dir", str(tmp_path / "cache"), "--out-dir", str(out),
            "--allow-partial", "--docx"]
    assert t.main(["go", site + "/sql/", "--yes", "--exclude", "missing|blocked|private|secret|captcha"] + base) == 0
    d = out / "tutorial-compilation"
    docs = list(d.glob("*.docx")) or list(out.glob("*/*.docx"))
    comp = json.loads(next(out.glob("*/completeness.json")).read_text())
    assert any(c["status"] == "skipped" for c in comp["chapters"])
    word = _Doc(str(docs[0]))
    assert not any("identical" in p.text or "missing from this document" in p.text for p in word.paragraphs)


def test_branding_and_user_agent():
    assert t.APP_NAME == "TutorialStudyForge" and t.APP_SHORT == "TSF" and "TutorialStudyForge-TSF" in t.USER_AGENT
    assert t.Settings().cache_dir.name == "tsf"
